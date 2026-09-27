# -*- coding: utf-8 -*-
"""
Скрипт восстановления эталонных ФИО и Должностей в файлах Excel и CSV.
Объединяет сохраняемые табельные номера, реальные ФИО и реальные должности.
"""

import openpyxl
import csv
from openpyxl.styles import Font, PatternFill, Alignment
from pathlib import Path

# Эталонный список реальных данных сотрудников
REAL_EMPLOYEES = [
    {"tab": "4004236", "fio": "Цюрко Геннадий Васильевич", "pos": "Слесарь по ремонту подвижного состава"},
    {"tab": "4604571", "fio": "Ханин Дмитрий Викторович", "pos": "Слесарь по ремонту подвижного состава"},
    {"tab": "4601556", "fio": "Тимин Игорь Витальевич", "pos": "Слесарь по ремонту подвижного состава"},
    {"tab": "4607514", "fio": "Будзинский Никита Андреевич", "pos": "Слесарь по ремонту подвижного состава"},
    {"tab": "4611025", "fio": "Александров Вадим Николаевич", "pos": "Слесарь по ремонту подвижного состава"},
    {"tab": "4016004", "fio": "Степанов Илья Борисович", "pos": "Слесарь-электрик по ремонту электрооборудования"},
    {"tab": "4610209", "fio": "Сивов Александр Сергеевич", "pos": "Слесарь-электрик по ремонту электрооборудования"},
    {"tab": "4004080", "fio": "Комлев Дмитрий Александрович", "pos": "Аккумуляторщик"},
    {"tab": "4610497", "fio": "Спиридонов Константин Герольдович", "pos": "Токарь"},
    {"tab": "4004267", "fio": "Щербакова Ирина Станиславовна", "pos": "Машинист крана"},
    {"tab": "4015690", "fio": "Цюрко Ирина Ивановна", "pos": "Машинист по стирке спецодежды"},
    {"tab": "4015696", "fio": "Шварчков Сергей Александрович", "pos": "Начальник участка"}
]

def main():
    xlsx_path = Path("employees_private.xlsx")
    csv_path = Path("employees_private.csv")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Сотрудники"

    # Заголовок таблицы
    headers = ["Код системы (ID)", "Должность", "ФИО", "Табельный номер"]
    ws.append(headers)

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    rows = []
    for i, emp in enumerate(REAL_EMPLOYEES, start=1):
        anon_id = f"ID_{i:03d}"
        row = [anon_id, emp["pos"], emp["fio"], emp["tab"]]
        ws.append(row)
        rows.append(row)

    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 45
    ws.column_dimensions['C'].width = 40
    ws.column_dimensions['D'].width = 20

    wb.save(xlsx_path)
    print(f"Файл {xlsx_path} полностью восстановлен с оригинальными ФИО!")

    # Сохраняем также CSV для быстрого открытия в Excel
    with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow(headers)
        writer.writerows(rows)

    print(f"Файл {csv_path} полностью восстановлен с оригинальными ФИО!")

if __name__ == "__main__":
    main()
