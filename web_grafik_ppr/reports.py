from __future__ import annotations

import calendar
import datetime as dt
import json
from io import BytesIO
from pathlib import Path
from urllib.parse import quote

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import sys
_PKG_ROOT = Path(__file__).resolve().parent
if str(_PKG_ROOT) not in sys.path:
    sys.path.insert(0, str(_PKG_ROOT))
if str(_PKG_ROOT.parent) not in sys.path:
    sys.path.insert(0, str(_PKG_ROOT.parent))

from rtps_common import connect_sqlite

try:
    from .constants import (
        ROOT,
        SOURCE_DIR,
        REPORT_TEMPLATE_NAME,
        TU28_TEMPLATE_NAME,
        ACT_TEMPLATE_NAME,
        MONTHS_RU,
        TEP_REPORT_FACTORS,
        AGR_REPORT_FACTORS,
        TU28_REPAIR_CODES,
    )
    from .calculations import (
        s,
        format_n,
        month_index,
        calculate_report_data_from_state,
        build_report_excel_tags,
        format_fio_initials,
        normalize_repair_code,
    )
    from .storage import (
        load_state,
        get_act_inventory_item,
        conn,
        DB_LOCK,
    )
except (ImportError, ValueError):
    from constants import (
        ROOT,
        SOURCE_DIR,
        REPORT_TEMPLATE_NAME,
        TU28_TEMPLATE_NAME,
        ACT_TEMPLATE_NAME,
        MONTHS_RU,
        TEP_REPORT_FACTORS,
        AGR_REPORT_FACTORS,
        TU28_REPAIR_CODES,
    )
    from calculations import (
        s,
        format_n,
        month_index,
        calculate_report_data_from_state,
        build_report_excel_tags,
        format_fio_initials,
        normalize_repair_code,
    )
    from storage import (
        load_state,
        get_act_inventory_item,
        conn,
        DB_LOCK,
    )

# Модуль формирования Excel отчетов (Отчет, Акт, ТУ-28) через openpyxl

def find_act_template_path() -> Path | None:
    candidates = [
        ROOT / ACT_TEMPLATE_NAME,
        SOURCE_DIR / ACT_TEMPLATE_NAME,
        ROOT.parent / "dist" / "РТПС" / "_internal" / "График ППР" / ACT_TEMPLATE_NAME,
    ]
    for candidate in candidates:
        if candidate.exists():
            return candidate
    return None


def find_report_template_path() -> Path | None:
    candidates = [
        ROOT / REPORT_TEMPLATE_NAME,
        SOURCE_DIR / REPORT_TEMPLATE_NAME,
        ROOT.parent / "dist" / "РТПС" / "_internal" / "График ППР" / REPORT_TEMPLATE_NAME,
    ]
    for candidate in candidates:
        if candidate.exists():
            return candidate
    return None


def find_tu28_template_path() -> Path | None:
    candidates = [
        ROOT / TU28_TEMPLATE_NAME,
        SOURCE_DIR / TU28_TEMPLATE_NAME,
        ROOT.parent / "dist" / "РТПС" / "_internal" / "График ППР" / TU28_TEMPLATE_NAME,
    ]
    for candidate in candidates:
        if candidate.exists():
            return candidate
    return None


def build_report_workbook(year: int, month_name: str, state: dict | None = None) -> tuple[bytes, str]:
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Alignment, Font
    except ImportError as exc:
        raise RuntimeError("На сервере не установлен openpyxl") from exc

    state = state or load_state(year)
    data = calculate_report_data_from_state(state, month_name)
    saved_notes = state.get("notes", {}).get(month_name, {}) or {}
    tags = build_report_excel_tags(month_name, data, saved_notes)

    template_path = find_report_template_path()
    if template_path:
        wb = load_workbook(template_path)
        replace_tags_in_workbook(wb, tags)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Отчет"
        for idx, (k, v) in enumerate(tags.items(), start=1):
            ws[f"A{idx}"] = k
            ws[f"B{idx}"] = v
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(horizontal="left", vertical="top")

    out = BytesIO()
    wb.save(out)
    return out.getvalue(), f"Отчет_{month_name}_{year}.xlsx"


def replace_tags_in_workbook(wb, tags: dict[str, str]) -> None:
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value:
                    value = cell.value
                    for tag, rep in tags.items():
                        if tag in value:
                            value = value.replace(tag, s(rep))
                    cell.value = value


def build_act_workbook(year: int, act: str) -> tuple[bytes, str]:
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Alignment, Font
    except ImportError as exc:
        raise RuntimeError("На сервере не установлен openpyxl") from exc

    clean_act_num = act.replace("Акт № ", "").strip()
    parts = clean_act_num.split("-")
    if len(parts) != 3:
        raise ValueError("Не удалось распознать формат акта")

    d_act, m_act, num_act = parts
    months_ru = {
        "01": "января", "02": "февраля", "03": "марта",
        "04": "апреля", "05": "мая", "06": "июня",
        "07": "июля", "08": "августа", "09": "сентября",
        "10": "октября", "11": "ноября", "12": "декабря",
    }
    date_str = f"{d_act} {months_ru.get(m_act, 'января')} {year} г."
    ser, inv = get_act_inventory_item(year, num_act)
    if "ПЭ" in ser.upper():
        eq_type = "Тяговый агрегат"
    elif ser:
        eq_type = "Тепловоз маневровый"
    else:
        eq_type = ""

    tags = {
        "[АКТ]": clean_act_num, "[Акт]": clean_act_num, "[акт]": clean_act_num,
        "[ДАТА]": date_str, "[Дата]": date_str,
        "[НОМЕР]": num_act, "[Номер]": num_act,
        "[АГРЕГАТ]": eq_type, "[Агрегат]": eq_type,
        "[СЕРИЯ]": ser, "[Серия]": ser,
        "[ИНВ]": inv, "[Инв]": inv,
    }

    template_path = find_act_template_path()
    if template_path:
        wb = load_workbook(template_path)
        replace_tags_in_workbook(wb, tags)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Акт"
        ws["A1"] = "Акт"
        ws["B1"] = clean_act_num
        ws["A2"] = "Дата"
        ws["B2"] = date_str
        ws["A3"] = "Серия"
        ws["B3"] = ser
        ws["A4"] = "Инвентарный номер"
        ws["B4"] = inv
        ws["A5"] = "Тип"
        ws["B5"] = eq_type
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for cell in ws["A1:B5"]:
            for item in cell:
                item.alignment = Alignment(horizontal="left")

    out = BytesIO()
    wb.save(out)
    return out.getvalue(), f"Акт_{clean_act_num}.xlsx"


def build_tu28_workbook(year: int, month_name: str, row_idx: int, staff_list: list[str] | None = None, state: dict | None = None, extra_repairs: list[str] | None = None) -> tuple[bytes, str]:
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Alignment, Font
    except ImportError as exc:
        raise RuntimeError("На сервере не установлен openpyxl") from exc

    state = state or load_state(year)
    month = next((m for m in state.get("months", []) if s(m.get("name")) == month_name), None)
    if not month:
        raise ValueError("Не удалось найти месяц")
    fact_rows = month.get("fact") or []
    if row_idx < 0 or row_idx >= len(fact_rows):
        raise ValueError("Не удалось найти строку ремонта")
    row = fact_rows[row_idx]
    cells = row.get("cells") or []
    series = s(cells[1]).strip()
    number = s(cells[2]).strip()
    if not number:
        raise ValueError("Не удалось определить номер")

    repair_code = ""
    repair_days = []
    for col in range(4, 4 + int(month.get("days") or 0)):
        value = normalize_repair_code(s(cells[col]) if col < len(cells) else "")
        if value in TU28_REPAIR_CODES:
            if not repair_code:
                repair_code = s(cells[col]).strip().upper()
            repair_days.append(col - 3)
    if not repair_days:
        raise ValueError("В выбранной строке не найден ремонт для ТУ-28")

    if "ПЭ" in series.upper():
        eq_type = "Тяговый агрегат"
    elif series:
        eq_type = "Тепловоз маневровый"
    else:
        eq_type = ""

    try:
        month_num = int(month.get("month") or 0)
    except Exception:
        month_num = 0
    repair_day = repair_days[0]
    date_str = f"{repair_day:02d}.{month_num:02d}.{year}"

    start_date_str = f"{year}-{month_num:02d}-{repair_days[0]:02d}"
    end_date_str = f"{year}-{month_num:02d}-{repair_days[-1]:02d}"

    db_path = Path(__file__).resolve().parent.parent / "base" / "common_database.db"
    if db_path.exists() and number and start_date_str and end_date_str:
        try:
            with connect_sqlite(db_path) as conn:
                cur = conn.cursor()
                db_rows = cur.execute(
                    """
                    SELECT measurement_date, r, c, v 
                    FROM archive_data 
                    WHERE locomotive=? AND measurement_date >= ? AND measurement_date <= ?
                    ORDER BY measurement_date DESC, r, c
                    """,
                    (number, start_date_str, end_date_str)
                ).fetchall()
                dates_dict = {}
                for d_str, r, c, v in db_rows:
                    if d_str not in dates_dict:
                        dates_dict[d_str] = {}
                    try:
                        r_int = int(r)
                        c_int = int(c)
                    except Exception:
                        continue
                    dates_dict[d_str].setdefault(r_int, {})[c_int] = str(v).strip() if v else ""
                if dates_dict:
                    best_date = sorted(dates_dict.keys(), reverse=True)[0]
                    measurements = dates_dict[best_date]
                    print(f"DEBUG: ZAMER KP: best_date={best_date} measurements dict: {measurements}", flush=True)
                    print(f"DEBUG: ZAMER KP: best_date={best_date} measurements size={len(measurements)}", flush=True)
        except Exception as e:
            print("Error reading Zamer KP archive db:", e, flush=True)
    tags = {
        "[СЕРИЯ]": series,
        "[НОМЕР]": number,
        "[ДАТА]": date_str,
        "[ВИД]": repair_code,
        "[АГРЕГАТ]": eq_type,
        "[ДИЗЕЛЬ]": "",
        "[ЭКИПАЖ 1]": "",
        "[ЭКИПАЖ 2]": "",
        "[АКБ]": "",
        "[ЭЛМАШ]": "",
        "[ЭЛАП]": "",
        "[ТОРМОЗ]": "",
    }

    col_to_tag_prefix = {
        2: "ПР_Л", 3: "ПР_П",
        4: "ТГ_Л", 5: "ТГ_П",
        6: "КР_Л", 7: "КР_П",
        8: "ТБ_Л", 9: "ТБ_П",
        10: "ДБ_Л", 11: "ДБ_П"
    }
    for axle in range(1, 13):
        for c_idx, prefix in col_to_tag_prefix.items():
            # In archive_data, r=2 corresponds to axle 1, r=3 to axle 2, etc.
            measurement_value = s(measurements.get(axle + 1, {}).get(c_idx, "")).strip()
            tags[f"[{prefix}_{axle}]"] = measurement_value.replace(".", ",")
            
    print(f"DEBUG: ZAMER KP TAGS: { {k: tags[k] for k in tags if 'ПР' in k or 'ТГ' in k} }", flush=True)

    components = [
        "ДИЗЕЛЬ",
        "ЭКИПАЖ 1",
        "ЭКИПАЖ 2",
        "АКБ",
        "ЭЛМАШ",
        "ЭЛАП",
        "ТОРМОЗ",
    ]
    for idx, name in enumerate(staff_list or []):
        if idx >= len(components):
            break
        tags[f"[{components[idx]}]"] = format_fio_initials(name)

    extra_repairs = extra_repairs or []
    print(f"DEBUG: build_tu28_workbook received extra_repairs={extra_repairs}", flush=True)
    for i in range(1, 21):
        tags[f"[ДОП_РЕМОНТ_{i}]"] = ""
        tags[f"[ДОП_НОМЕР_{i}]"] = ""
    for i, extra in enumerate(extra_repairs, start=1):
        if i <= 20:
            val = str(extra).strip()
            if val:
                tags[f"[ДОП_РЕМОНТ_{i}]"] = val
                tags[f"[ДОП_НОМЕР_{i}]"] = str(i)
    print(f"DEBUG: tags built: {[k for k in tags.keys() if 'ДОП' in k and tags[k]]}", flush=True)

    template_path = find_tu28_template_path()
    if template_path:
        wb = load_workbook(template_path)
        replace_tags_in_workbook(wb, tags)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "ТУ-28"
        for idx, (k, v) in enumerate(tags.items(), start=1):
            ws[f"A{idx}"] = k
            ws[f"B{idx}"] = v
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row_cells in ws.iter_rows():
            for cell in row_cells:
                cell.alignment = Alignment(horizontal="left", vertical="top")

    out = BytesIO()
    wb.save(out)
    return out.getvalue(), f"ТУ-28_{number}_{month_name}_{year}.xlsx"


def content_disposition_attachment(filename: str) -> str:
    ascii_name = filename.encode("ascii", "ignore").decode("ascii") or "file.xlsx"
    return f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{quote(filename)}"

