from __future__ import annotations

import base64
import calendar
import copy
import datetime as dt
import io
import re
import tempfile
import sys
from pathlib import Path

# Обеспечиваем доступность путей модуля и корня проекта для импортов
_ROOT = Path(__file__).resolve().parent
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))
if str(_ROOT.parent) not in sys.path:
    sys.path.insert(0, str(_ROOT.parent))

import openpyxl

# Импорт констант и функций доступа к данным
try:
    from .constants import (
        DB_LOCK,
        MONTH_NAMES,
        ROOT,
    )
    from .storage import (
        conn,
        load_system_dates,
        text,
    )
except (ImportError, ValueError):
    from constants import (
        DB_LOCK,
        MONTH_NAMES,
        ROOT,
    )
    from storage import (
        conn,
        load_system_dates,
        text,
    )


# Клиентский скрипт автоматической расшифровки меток сотрудников (Работник №X -> реальное ФИО) из localStorage
HTML_DEANONYMIZE_SCRIPT = """
<script>
(function() {
  function getEmpReplacementPairs() {
    const raw = localStorage.getItem('rtps_emp_dict') || localStorage.getItem('rtps_employees_full_dict');
    if (!raw) return [];
    let map = null;
    try { map = JSON.parse(raw); } catch (e) { return []; }
    if (!map) return [];
    const flatMap = {};
    function extractNum(idStr) {
      if (!idStr) return null;
      const clean = String(idStr).replace(/ID_/g, '').replace(/Работник\\s*№?\\s*/g, '').replace(/СотрудникПолн\\s*№?\\s*/g, '').replace(/Сотрудник\\s*№?\\s*/g, '').replace(/Должность\\s*№?\\s*/g, '').trim();
      const num = parseInt(clean, 10);
      return isNaN(num) ? null : num;
    }
    function addFieldVariants(prefix, num, val) {
      const sNum = String(num);
      const pNum = (num < 10 ? '00' : (num < 100 ? '0' : '')) + num;
      const valStr = String(val);
      flatMap[prefix + " №" + num] = valStr;
      flatMap[prefix + " №" + sNum] = valStr;
      flatMap[prefix + " №" + pNum] = valStr;
      flatMap[prefix + " № " + sNum] = valStr;
      flatMap[prefix + " № " + pNum] = valStr;
    }
    function addIdVariants(num, val) {
      const sNum = String(num);
      const pNum = (num < 10 ? '00' : (num < 100 ? '0' : '')) + num;
      const valStr = String(val);
      flatMap["ID_" + pNum] = valStr;
      flatMap["ID_" + sNum] = valStr;
      flatMap["ID_" + num] = valStr;
    }
    if (Array.isArray(map)) {
      for (let i = 0; i < map.length; i++) {
        let rec = map[i];
        if (!rec) continue;
        let num = extractNum(rec.id || rec.tab || rec.tab_num);
        if (num === null) num = i + 1;
        let shortF = rec.shortFio || rec.fio || rec.name || rec.short_name || '';
        let fullF = rec.fullFio || rec.full_name || rec.fio || rec.name || '';
        let pos = rec.pos || rec.position || '';
        let tab = rec.tab || rec.tab_num || rec.id || '';
        if (shortF) addFieldVariants("Работник", num, shortF);
        if (fullF) {
          addFieldVariants("СотрудникПолн", num, fullF);
          addFieldVariants("Сотрудник", num, fullF);
        }
        if (pos) addFieldVariants("Должность", num, pos);
        if (tab) {
          addIdVariants(num, tab);
          if (rec.id) flatMap[String(rec.id)] = String(tab);
        }
      }
    } else if (typeof map === 'object') {
      for (let k in map) {
        if (map.hasOwnProperty(k) && k && map[k]) {
          const valStr = String(map[k]);
          flatMap[k] = valStr;
          let num = extractNum(k);
          if (num !== null) {
            if (k.startsWith("Работник")) addFieldVariants("Работник", num, valStr);
            else if (k.startsWith("СотрудникПолн") || k.startsWith("Сотрудник")) {
              addFieldVariants("СотрудникПолн", num, valStr);
              addFieldVariants("Сотрудник", num, valStr);
            } else if (k.startsWith("Должность")) addFieldVariants("Должность", num, valStr);
            else if (k.startsWith("ID_") || /^\\d+$/.test(k)) addIdVariants(num, valStr);
          }
        }
      }
    }
    const pairs = [];
    for (let k in flatMap) {
      if (flatMap.hasOwnProperty(k) && k && flatMap[k]) {
        pairs.push({ from: k, to: String(flatMap[k]) });
      }
    }
    pairs.sort((a, b) => b.from.length - a.from.length);
    return pairs;
  }

  function deAnonymize() {
    const pairs = getEmpReplacementPairs();
    if (!pairs || !pairs.length) return;

    // Замена текстовых узлов во всем документе
    const walker = document.createTreeWalker(document.body, NodeFilter.SHOW_TEXT, null, false);
    let node;
    while ((node = walker.nextNode())) {
      let val = node.nodeValue;
      if (!val || !val.trim()) continue;
      let newVal = val;
      for (let p = 0; p < pairs.length; p++) {
        if (newVal.indexOf(pairs[p].from) !== -1) {
          newVal = newVal.split(pairs[p].from).join(pairs[p].to);
        }
      }
      if (newVal !== val) node.nodeValue = newVal;
    }

    // Сортировка строк таблицы по реальному ФИО в алфавитном порядке
    const table = document.querySelector('table');
    if (table) {
      const rows = Array.from(table.querySelectorAll('tr')).slice(1);
      const totalRow = rows.length > 0 && rows[rows.length - 1].textContent.includes('Итого') ? rows.pop() : null;
      if (rows.length > 0) {
        // Определяем индекс колонки с ФИО (по умолчанию 1)
        const headers = Array.from(table.querySelectorAll('th'));
        let fioColIdx = 1;
        headers.forEach((th, idx) => {
          if (th.textContent.includes('ФИО') || th.textContent.includes('Сотрудник')) {
            fioColIdx = idx;
          }
        });

        rows.sort((a, b) => {
          const nameA = a.children[fioColIdx] ? a.children[fioColIdx].textContent.trim() : '';
          const nameB = b.children[fioColIdx] ? b.children[fioColIdx].textContent.trim() : '';
          return nameA.localeCompare(nameB, 'ru');
        });
        rows.forEach((r, idx) => {
          if (r.children[0] && (r.children[0].classList.contains('center') || r.children[0].classList.contains('num'))) {
            r.children[0].textContent = idx + 1;
          }
          table.appendChild(r);
        });
        if (totalRow) table.appendChild(totalRow);
      }
    }
  }

  if (document.readyState === 'loading') {
    document.addEventListener('DOMContentLoaded', deAnonymize);
  } else {
    deAnonymize();
  }
})();
</script>
"""


def build_summary_html(year: int, month: int, report_type: str) -> str:
    """Генерация сводного HTML-отчета по отпускам, больничным или работе в выходные дни."""
    if ":" in report_type:
        codes = [report_type.split(":")[0].strip()]
        title = report_type.split(":", 1)[1].strip()
    else:
        code_map = {
            "Отпуска": ["О", "ОВ"],
            "Отпуска внеплановые": ["ОВ"],
            "Отпуск б/с": ["ДО"],
            "Учебный отпуск": ["У"],
            "Больничный": ["Б"],
        }
        codes = code_map.get(report_type, [report_type])
        title = report_type

    with DB_LOCK, conn() as connection:
        cur = connection.cursor()

        emp_rows = cur.execute("SELECT DISTINCT name FROM employees WHERE y=? AND name != ''", (year,)).fetchall()
        employees = sorted(list(set(row["name"] for row in emp_rows)))
        result = {emp: 0 for emp in employees}

        sys_dates = load_system_dates(year)
        holiday_set = set(sys_dates.get("holiday", []))
        transfer_set = set(sys_dates.get("transfer", []))

        months_names = ["Январь", "Февраль", "Март", "Апрель", "Май", "Июнь", "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"]
        month_map = {name: i + 1 for i, name in enumerate(months_names)}

        if codes == ["WORK_WEEKEND"]:
            ts_rows = cur.execute(
                """
                SELECT e.name, t.v, t.m, t.c
                FROM timesheet t
                JOIN employees e ON t.tab_num = e.tab_num AND t.y = e.y
                WHERE t.y=? AND t.v != ''
                """,
                (year,),
            ).fetchall()

            ignore_codes = {"В", "B", "О", "ДО", "К", "У", "Б", "БН", "ОВ"}

            for row in ts_rows:
                name = row["name"]
                v = str(row["v"]).strip().upper()
                if not v or v in ignore_codes:
                    continue

                m_str = row["m"]
                d = int(row["c"])
                m_int = month_map.get(m_str)
                if not m_int:
                    continue

                is_holiday = (m_int, d) in holiday_set
                is_transfer = (m_int, d) in transfer_set
                try:
                    weekday = dt.date(year, m_int, d).weekday()
                    is_weekend = weekday in (5, 6)
                except ValueError:
                    is_weekend = False

                if is_weekend or is_holiday or is_transfer:
                    if name in result:
                        result[name] += 1
        else:
            placeholders = ",".join(["?"] * len(codes))
            query = f"""
                SELECT e.name, t.v, t.m, t.c
                FROM timesheet t
                JOIN employees e ON t.tab_num = e.tab_num AND t.y = e.y
                WHERE t.y=? AND t.v IN ({placeholders})
            """
            ts_rows = cur.execute(query, [year] + codes).fetchall()

            vacation_days = {}
            if "О" in codes or "ОВ" in codes:
                vac_rows = cur.execute(
                    "SELECT e.name, v.c, v.v as val FROM vacations v JOIN employees e ON v.tab_num = e.tab_num AND v.y = e.y WHERE v.y=?",
                    (year,),
                ).fetchall()
                emp_vacs = {}
                for r in vac_rows:
                    emp_vacs.setdefault(r["name"], {})[int(r["c"])] = r["val"]

                def parse_vacation_date(date_str, default_year):
                    if not date_str:
                        return None
                    s = str(date_str).strip()
                    if "." in s:
                        p = s.split(".")
                        try:
                            if len(p) == 2:
                                return dt.date(default_year, int(p[1]), int(p[0]))
                            elif len(p) == 3:
                                y = int(p[2])
                                if y < 100:
                                    y += 2000
                                return dt.date(y, int(p[1]), int(p[0]))
                        except Exception:
                            pass
                    return None

                for name, vac_data in emp_vacs.items():
                    vacation_days[name] = set()
                    for pair in [(1, 2), (5, 6), (9, 10)]:
                        s_str = vac_data.get(pair[0])
                        e_str = vac_data.get(pair[1])
                        try:
                            s_date = parse_vacation_date(s_str, year)
                            e_date = parse_vacation_date(e_str, year)
                            if s_date and e_date and e_date >= s_date:
                                curr = s_date
                                while curr <= e_date:
                                    if curr.year == year:
                                        vacation_days[name].add((curr.month, curr.day))
                                    curr += dt.timedelta(days=1)
                        except Exception:
                            pass

            for row in ts_rows:
                name = row["name"]
                v = row["v"]
                m_str = row["m"]
                d = row["c"]

                if v in ("О", "ОВ"):
                    m_int = month_map.get(m_str)
                    if m_int:
                        if name not in vacation_days:
                            vacation_days[name] = set()
                        vacation_days[name].add((m_int, int(d)))
                else:
                    if name in result:
                        result[name] += 1

            if "О" in codes or "ОВ" in codes:
                for name, days_set in vacation_days.items():
                    valid_days = 0
                    for m_int, d_int in days_set:
                        if (m_int, d_int) not in holiday_set:
                            valid_days += 1
                    if name in result:
                        result[name] += valid_days

    total = sum(result.values())

    html = f"""
    <html><head><meta charset="utf-8"><title>Сводка: {title}</title>
    <style>
        body {{ font-family: Arial, sans-serif; margin: 20px; }}
        table {{ border-collapse: collapse; width: 600px; margin-top: 20px; }}
        th, td {{ border: 1px solid #ccc; padding: 8px; text-align: left; }}
        th {{ background-color: #f0f0f0; }}
        .center {{ text-align: center; }}
        .right {{ text-align: right; font-weight: bold; }}
    </style>
    </head><body>
    <h2 class="center">{title} за {year} год</h2>
    <div class="center" style="color: #666;">Коды: {', '.join(codes)}</div>
    <table>
        <tr><th style="width: 50px;">№</th><th>ФИО</th><th style="width: 100px;">Дней</th></tr>
    """
    for idx, emp in enumerate(employees, 1):
        html += f"<tr><td class='center'>{idx}</td><td>{emp}</td><td class='center'>{result[emp]}</td></tr>"

    html += f"<tr><td colspan='2' class='right'>Итого:</td><td class='center'><b>{total}</b></td></tr>"
    html += f"</table>{HTML_DEANONYMIZE_SCRIPT}</body></html>"
    return html


def build_milk_report_excel(year: int, month: int, report_type: str) -> tuple[str, str]:
    """Генерация ведомости на выдачу молока или компенсации в формате Excel (.xlsx)."""
    templates = {
        "компенсация": "Молоко_комп_шаблон.xlsx",
        "план": "Молоко_план_шаблон.xlsx",
        "факт": "Молоко_факт_шаблон.xlsx",
    }

    template_name = templates.get(report_type)
    if not template_name:
        raise ValueError(f"Неизвестный тип отчета по молоку: {report_type}")

    template_path = ROOT / "resources" / template_name
    if not template_path.exists():
        raise FileNotFoundError(f"Файл шаблона не найден: {template_name}")

    sys_dates = load_system_dates(year)
    transfer_dates = sys_dates["transfer"]
    holiday_dates = sys_dates["holiday"]

    m_str = MONTH_NAMES[month] if 1 <= month <= 12 else str(month)
    days_cnt = calendar.monthrange(year, month)[1]

    def employee_exclude_start(raw_date: str) -> int | None:
        raw = text(raw_date).strip()
        if not raw:
            return None
        parts = raw.split("-")
        if len(parts) != 3:
            return None
        try:
            ex_year = int(parts[0])
            ex_month = int(parts[1])
            ex_day = int(parts[2])
        except Exception:
            return None
        if year > ex_year:
            return 1
        if year == ex_year and month > ex_month:
            return 1
        if year == ex_year and month == ex_month:
            return ex_day
        return None

    def employee_hire_start(raw_date: str) -> int | None:
        raw = text(raw_date).strip()
        if not raw:
            return None
        parts = raw.split("-")
        if len(parts) != 3:
            return None
        try:
            h_year = int(parts[0])
            h_month = int(parts[1])
            h_day = int(parts[2])
        except Exception:
            return None
        if year < h_year:
            return 32
        if year == h_year and month < h_month:
            return 32
        if year == h_year and month == h_month:
            return h_day
        return None

    def is_workday(y, m, d):
        dt_obj = dt.date(y, m, d)
        is_we = dt_obj.weekday() >= 5
        is_hol = (m, d) in holiday_dates
        is_tr = (m, d) in transfer_dates
        if is_tr:
            return True
        if is_hol:
            return False
        return not is_we

    non_shift_codes = {"В", "B", "О", "ОВ", "А", "У", "Б", "БН"}
    numeric_shift_re = re.compile(r"^\d+(?:[.,]\d+)?$")
    spacing_re = re.compile(r"\s*,\s*")

    def is_shift_mark(value: str) -> bool:
        val = text(value).strip().upper()
        if not val:
            return False
        compact = spacing_re.sub(",", val).replace(" ", "")
        if "М" in compact:
            return True
        if compact in non_shift_codes:
            return False
        return bool(numeric_shift_re.match(compact))

    def build_report_rows() -> tuple[list[dict], int]:
        final_rows: list[dict] = []
        grand_total = 0

        with DB_LOCK, conn() as connection:
            cur = connection.cursor()

            emp_rows = cur.execute(
                "SELECT pos, name, tab_num, milk, milk_issue, full_name, milk_note, hire_date, exclude_date FROM employees WHERE y=? AND name != '' ORDER BY rowid",
                (year,),
            ).fetchall()

            m_comp_set = set()
            m_issue_set = set()

            for r in emp_rows:
                name = str(r["name"]).upper()
                if r["milk"]:
                    m_comp_set.add(name)
                if r["milk_issue"]:
                    m_issue_set.add(name)

            ts_rows = cur.execute("SELECT tab_num, c, v FROM timesheet WHERE y=? AND m=?", (year, m_str)).fetchall()
            ts_data = {}
            for r in ts_rows:
                ts_data.setdefault(str(r["tab_num"]), {})[int(r["c"])] = str(r["v"]).upper()

            def allowed_employee(emp_row):
                name_up = str(emp_row["name"]).upper()
                if report_type == "компенсация":
                    return name_up in m_comp_set
                return name_up in m_issue_set

            for emp in emp_rows:
                if not allowed_employee(emp):
                    continue

                raw_tab = str(emp["tab_num"])
                digits = "".join(ch for ch in raw_tab if ch.isdigit())
                num_part = int(digits) if digits else 1

                name = f"Работник №{num_part}"
                tab_num = raw_tab
                pos = f"Должность №{num_part}"
                full_name = f"СотрудникПолн №{num_part}"
                milk_note = str(emp["milk_note"])
                exclude_start = employee_exclude_start(emp["exclude_date"] if "exclude_date" in emp.keys() else "")
                hire_start = employee_hire_start(emp["hire_date"] if "hire_date" in emp.keys() else "")

                if exclude_start == 1 or hire_start == 32:
                    continue

                count = 0
                missed_days: list[str] = []
                emp_ts = ts_data.get(tab_num, {})

                for d in range(1, days_cnt + 1):
                    if exclude_start is not None and d >= exclude_start:
                        continue
                    if hire_start is not None and d < hire_start:
                        continue

                    val = emp_ts.get(d, "").strip().upper()
                    if report_type == "компенсация":
                        if "М" in val or "M" in val:
                            count += 1
                        elif val:
                            missed_days.append(f"{d:02d}.{month:02d} {val}")
                    elif report_type == "факт":
                        if is_shift_mark(val) and val not in non_shift_codes:
                            count += 1
                        elif val:
                            missed_days.append(f"{d:02d}.{month:02d} {val}")
                    else:
                        if not is_workday(year, month, d):
                            continue
                        if not val:
                            count += 1
                        elif val not in non_shift_codes and is_shift_mark(val):
                            count += 1
                        else:
                            missed_days.append(f"{d:02d}.{month:02d} {val or 'пусто'}")

                final_rows.append(
                    {
                        "fio": name,
                        "full_name": full_name,
                        "pos": pos,
                        "tab": tab_num,
                        "shifts": count,
                        "missed_note": "; ".join(missed_days),
                        "milk_note": milk_note,
                    }
                )
                grand_total += count

        return final_rows, grand_total

    with DB_LOCK, conn() as connection:
        cur = connection.cursor()
        norm_row = cur.execute("SELECT v FROM ts_norms_data WHERE y=? AND r=? AND c=2", (year, month - 1)).fetchone()
        work_days_norm = str(norm_row["v"]) if norm_row else "0"

    final_list, grand_total = build_report_rows()

    wb = openpyxl.load_workbook(template_path)
    ws = wb.active

    row_tpl = None
    for row in ws.iter_rows():
        for cell in row:
            if cell.value and isinstance(cell.value, str):
                cell.value = (
                    cell.value.replace("[МЕСЯЦ]", m_str)
                    .replace("[ГОД]", str(year))
                    .replace("[НОРМА_ДНЕЙ]", work_days_norm)
                    .replace("[ИТОГО]", str(grand_total))
                )
                if row_tpl is None and any(
                    tag in cell.value
                    for tag in [
                        "[№]",
                        "[ФИО]",
                        "[ФИО_ПОЛНОЕ]",
                        "[ДОЛЖНОСТЬ]",
                        "[ТАБ]",
                        "[СМЕНЫ]",
                        "[МОЛОКО_ПРИМ]",
                    ]
                ):
                    row_tpl = cell.row

    if row_tpl is not None and final_list:
        tpl_vals = {c: ws.cell(row=row_tpl, column=c).value for c in range(1, ws.max_column + 1)}
        ws.insert_rows(row_tpl + 1, len(final_list) - 1)
        for i, data in enumerate(final_list):
            curr_r = row_tpl + i
            ws.row_dimensions[curr_r].height = None
            for c_idx in range(1, ws.max_column + 1):
                cell = ws.cell(row=curr_r, column=c_idx)
                if i > 0:
                    src = ws.cell(row=row_tpl, column=c_idx)
                    cell.value = src.value
                    if src.has_style:
                        cell.font = copy.copy(src.font)
                        cell.border = copy.copy(src.border)
                        cell.fill = copy.copy(src.fill)
                        cell.alignment = copy.copy(src.alignment)
                v = tpl_vals.get(c_idx)
                if v and isinstance(v, str):
                    v = (
                        v.replace("[№]", str(i + 1))
                        .replace("[ФИО]", data["fio"])
                        .replace("[ФИО_ПОЛНОЕ]", data["full_name"])
                        .replace("[ДОЛЖНОСТЬ]", data["pos"])
                        .replace("[ТАБ]", data["tab"])
                        .replace("[СМЕНЫ]", str(data["shifts"]))
                        .replace("[МОЛОКО_ПРИМ]", data["milk_note"])
                    )
                    cell.value = int(v) if str(v).isdigit() else v

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    tmp.close()
    wb.save(tmp.name)

    filename_formatted = f"Отчет_Молоко_{report_type}_{m_str}_{year}.xlsx"
    return tmp.name, filename_formatted


def build_milk_details_html(year: int, month: int, report_type: str) -> str:
    """Генерация детального отчета по незасчитанным дням для проверки выдачи молока."""
    titles = {
        "компенсация": "Компенсация (План)",
        "план": "Выдача (План)",
        "факт": "Выдача (Факт)",
    }
    title = titles.get(report_type, report_type)

    sys_dates = load_system_dates(year)
    transfer_dates = sys_dates["transfer"]
    holiday_dates = sys_dates["holiday"]
    m_str = MONTH_NAMES[month] if 1 <= month <= 12 else str(month)
    days_cnt = calendar.monthrange(year, month)[1]

    def employee_exclude_start(raw_date: str) -> int | None:
        raw = text(raw_date).strip()
        if not raw:
            return None
        parts = raw.split("-")
        if len(parts) != 3:
            return None
        try:
            ex_year = int(parts[0])
            ex_month = int(parts[1])
            ex_day = int(parts[2])
        except Exception:
            return None
        if year > ex_year:
            return 1
        if year == ex_year and month > ex_month:
            return 1
        if year == ex_year and month == ex_month:
            return ex_day
        return None

    def employee_hire_start(raw_date: str) -> int | None:
        raw = text(raw_date).strip()
        if not raw:
            return None
        parts = raw.split("-")
        if len(parts) != 3:
            return None
        try:
            h_year = int(parts[0])
            h_month = int(parts[1])
            h_day = int(parts[2])
        except Exception:
            return None
        if year < h_year:
            return 32
        if year == h_year and month < h_month:
            return 32
        if year == h_year and month == h_month:
            return h_day
        return None

    def is_workday(y, m, d):
        dt_obj = dt.date(y, m, d)
        is_we = dt_obj.weekday() >= 5
        is_hol = (m, d) in holiday_dates
        is_tr = (m, d) in transfer_dates
        if is_tr:
            return True
        if is_hol:
            return False
        return not is_we

    non_shift_codes = {"В", "B", "О", "ОВ", "А", "У", "Б", "БН"}
    numeric_shift_re = re.compile(r"^\d+(?:[.,]\d+)?$")
    spacing_re = re.compile(r"\s*,\s*")

    def is_shift_mark(value: str) -> bool:
        val = text(value).strip().upper()
        if not val:
            return False
        compact = spacing_re.sub(",", val).replace(" ", "")
        if "М" in compact:
            return True
        if compact in non_shift_codes:
            return False
        return bool(numeric_shift_re.match(compact))

    with DB_LOCK, conn() as connection:
        cur = connection.cursor()
        emp_rows = cur.execute(
            "SELECT pos, name, tab_num, milk, milk_issue, full_name, milk_note, hire_date, exclude_date FROM employees WHERE y=? AND name != '' ORDER BY rowid",
            (year,),
        ).fetchall()

        m_comp_set = {str(r["name"]).upper() for r in emp_rows if r["milk"]}
        m_issue_set = {str(r["name"]).upper() for r in emp_rows if r["milk_issue"]}
        ts_rows = cur.execute("SELECT tab_num, c, v FROM timesheet WHERE y=? AND m=?", (year, m_str)).fetchall()
        ts_data = {}
        for r in ts_rows:
            ts_data.setdefault(str(r["tab_num"]), {})[int(r["c"])] = str(r["v"]).upper()

    rows = []
    for emp in emp_rows:
        name = str(emp["name"])
        name_up = name.upper()
        if report_type == "компенсация" and name_up not in m_comp_set:
            continue
        if report_type in ("план", "факт") and name_up not in m_issue_set:
            continue

        raw_tab = str(emp["tab_num"])
        digits = "".join(ch for ch in raw_tab if ch.isdigit())
        num_part = int(digits) if digits else 1

        tab_num = raw_tab
        name = f"Работник №{num_part}"
        pos = f"Должность №{num_part}"
        full_name = f"СотрудникПолн №{num_part}"
        exclude_start = employee_exclude_start(emp["exclude_date"] if "exclude_date" in emp.keys() else "")
        hire_start = employee_hire_start(emp["hire_date"] if "hire_date" in emp.keys() else "")

        if exclude_start == 1 or hire_start == 32:
            continue

        emp_ts = ts_data.get(tab_num, {})
        count = 0
        missed_days: list[str] = []

        for d in range(1, days_cnt + 1):
            if exclude_start is not None and d >= exclude_start:
                continue
            if hire_start is not None and d < hire_start:
                continue
            val = emp_ts.get(d, "").strip().upper()
            if report_type == "компенсация":
                if "М" in val or "M" in val:
                    count += 1
                elif val:
                    missed_days.append(f"{d:02d}.{month:02d} — {val}")
            elif report_type == "факт":
                if is_shift_mark(val) and val not in non_shift_codes:
                    count += 1
                elif val:
                    missed_days.append(f"{d:02d}.{month:02d} — {val}")
            else:
                if not is_workday(year, month, d):
                    continue
                if not val:
                    count += 1
                elif val not in non_shift_codes and is_shift_mark(val):
                    count += 1
                else:
                    missed_days.append(f"{d:02d}.{month:02d} — {val or 'пусто'}")

        rows.append(
            {
                "tab": tab_num,
                "fio": full_name or name,
                "pos": pos,
                "shifts": count,
                "missed": missed_days,
            }
        )

    html = [f"<html><head><meta charset='utf-8'><title>{title} — {MONTH_NAMES[month]} {year}</title>"]
    html.append(
        """
    <style>
      body { font-family: Arial, sans-serif; margin: 20px; color: #1f2937; }
      h1 { margin: 0 0 8px; font-size: 22px; }
      .sub { color: #6b7280; margin-bottom: 16px; }
      table { border-collapse: collapse; width: 100%; }
      th, td { border: 1px solid #cbd5e1; padding: 8px 10px; vertical-align: top; }
      th { background: #eef4ff; text-align: left; }
      tr:nth-child(even) td { background: #fafcff; }
      .num { text-align: center; white-space: nowrap; }
      .missed { white-space: pre-wrap; color: #b45309; }
      .empty { color: #94a3b8; }
    </style>
    </head><body>"""
    )
    html.append(f"<h1>{title} — {MONTH_NAMES[month]} {year}</h1>")
    html.append(
        "<div class='sub'>Показываем только те дни, которые не были засчитаны как смена. Если строка пустая, для этого сотрудника не найдено неучтённых дней.</div>"
    )
    html.append(
        "<table><thead><tr><th style='width:50px'>№</th><th>Таб. №</th><th>ФИО</th><th>Профессия</th><th style='width:90px'>Смен</th><th>Не засчитано</th></tr></thead><tbody>"
    )
    for idx, row in enumerate(rows, 1):
        missed = "<div class='empty'>нет</div>" if not row["missed"] else "<br>".join(row["missed"])
        html.append(
            f"<tr><td class='num'>{idx}</td><td class='num'>{row['tab']}</td><td>{row['fio']}</td><td>{row['pos']}</td><td class='num'>{row['shifts']}</td><td class='missed'>{missed}</td></tr>"
        )
    html.append(f"</tbody></table>{HTML_DEANONYMIZE_SCRIPT}</body></html>")
    return "".join(html)


def build_sick_email_bytes(emp: str, op_type: str, start: str, end: str, email: str) -> bytes:
    """Генерация файла черновика электронного письма .eml по больничному листу."""
    with DB_LOCK, conn() as connection:
        cur = connection.cursor()
        emp_data = cur.execute("SELECT tab_num FROM employees WHERE name=?", (emp,)).fetchone()
        tab_num = str(emp_data["tab_num"]) if emp_data else ""

    html_body = f"""
    <html>
    <head>
    <style>
        table {{ border-collapse: collapse; width: 550px; font-family: Arial, sans-serif; border: 1px solid #000000; }}
        td {{ border: 1px solid #000000; padding: 6px; text-align: left; font-size: 13px; }}
    </style>
    </head>
    <body>
        <table>
            <tr>
                <td style="width: 40%;">Табельный номер</td>
                <td style="width: 60%;">{tab_num}</td>
            </tr>
            <tr>
                <td>ФИО сотрудника</td>
                <td>{emp}</td>
            </tr>
            <tr>
                <td>Тип операции</td>
                <td>{op_type}</td>
            </tr>
            <tr>
                <td style="padding-left: 20px;">Дата начала</td>
                <td>{start}</td>
            </tr>
            <tr>
                <td style="padding-left: 20px;">Дата окончания</td>
                <td>{end}</td>
            </tr>
            <tr>
                <td>Структурное подразделение</td>
                <td>3040 Рудник Таймырский</td>
            </tr>
        </table>
        <br>
    </body>
    </html>
    """

    encoded_subject = base64.b64encode("3040".encode("utf-8")).decode("utf-8")
    mime_subject = f"=?utf-8?B?{encoded_subject}?="

    eml_headers = [
        "MIME-Version: 1.0",
        f"To: {email}",
        f"Subject: {mime_subject}",
        "X-Unsent: 1",
        "Content-Type: text/html; charset=utf-8",
        "Content-Transfer-Encoding: 8bit",
        "",
        html_body,
    ]

    eml_data = "\r\n".join(eml_headers)
    return eml_data.encode("utf-8")
