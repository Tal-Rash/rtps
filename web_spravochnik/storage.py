from __future__ import annotations

import datetime as dt
import sqlite3
import sys
from contextlib import contextmanager
from pathlib import Path
from typing import Generator

# Добавляем путь к корню для импорта общей библиотеки rtps_common
_ROOT = Path(__file__).resolve().parent
if str(_ROOT.parent) not in sys.path:
    sys.path.insert(0, str(_ROOT.parent))

from rtps_common import connect_sqlite

try:
    from .constants import DB_FILE, DB_LOCK, MONTHS, ROOT, text
except (ImportError, ValueError):
    from constants import DB_FILE, DB_LOCK, MONTHS, ROOT, text


@contextmanager
def conn_ctx() -> Generator[sqlite3.Connection, None, None]:
    """Безопасный контекстный менеджер SQLite соединения с оптимизированными настройками."""
    connection = connect_sqlite(DB_FILE)
    connection.row_factory = sqlite3.Row
    try:
        connection.execute("PRAGMA journal_mode=WAL")
        connection.execute("PRAGMA synchronous=NORMAL")
        connection.execute("PRAGMA busy_timeout=5000")
        yield connection
    finally:
        connection.close()


def ensure_db() -> None:
    """Инициализация схемы таблиц сотрудников, локомотивов и норм времени."""
    DB_FILE.parent.mkdir(parents=True, exist_ok=True)
    with conn_ctx() as conn:
        with conn:
            cur = conn.cursor()

            # Таблица сотрудников
            cur.execute("""
                CREATE TABLE IF NOT EXISTS employees (
                    y INT, pos TEXT, name TEXT, tab_num TEXT,
                    milk INT DEFAULT 0, milk_issue INT DEFAULT 0,
                    exclude_date TEXT DEFAULT '',
                    full_name TEXT DEFAULT '', milk_note TEXT DEFAULT '',
                    PRIMARY KEY(y,pos,name)
                )
            """)
            emp_cols = {row[1] for row in cur.execute("PRAGMA table_info(employees)").fetchall()}
            if "exclude_date" not in emp_cols:
                cur.execute("ALTER TABLE employees ADD COLUMN exclude_date TEXT DEFAULT ''")
            if "hire_date" not in emp_cols:
                cur.execute("ALTER TABLE employees ADD COLUMN hire_date TEXT DEFAULT ''")
            if "vacation_days" not in emp_cols:
                cur.execute("ALTER TABLE employees ADD COLUMN vacation_days INT DEFAULT 28")

            # Таблица инвентаря (парк локомотивов)
            cur.execute("""
                CREATE TABLE IF NOT EXISTS inventory (
                    y INT, ser TEXT, num TEXT, inv TEXT,
                    PRIMARY KEY(y,ser,num)
                )
            """)
            inv_cols = {row[1] for row in cur.execute("PRAGMA table_info(inventory)").fetchall()}
            if "sort_order" not in inv_cols:
                cur.execute("ALTER TABLE inventory ADD COLUMN sort_order INT")
            if "updated_at" not in inv_cols:
                cur.execute("ALTER TABLE inventory ADD COLUMN updated_at INT NOT NULL DEFAULT 0")
            if "deleted_at" not in inv_cols:
                cur.execute("ALTER TABLE inventory ADD COLUMN deleted_at INT NOT NULL DEFAULT 0")
            if "wheel_pair_count" not in inv_cols:
                cur.execute("ALTER TABLE inventory ADD COLUMN wheel_pair_count INT")
            if "section_count" not in inv_cols:
                cur.execute("ALTER TABLE inventory ADD COLUMN section_count INT")
            if "eight_digit_number" not in inv_cols:
                cur.execute("ALTER TABLE inventory ADD COLUMN eight_digit_number TEXT")
            if "manufacture_year" not in inv_cols:
                cur.execute("ALTER TABLE inventory ADD COLUMN manufacture_year TEXT")
            if "service_life" not in inv_cols:
                cur.execute("ALTER TABLE inventory ADD COLUMN service_life TEXT")

            cur.execute("UPDATE inventory SET sort_order = rowid WHERE sort_order IS NULL OR sort_order <= 0")

            # Таблица норм времени (ТС)
            cur.execute("""
                CREATE TABLE IF NOT EXISTS ts_norms_data (
                    y INT, r INT, c INT, v TEXT,
                    PRIMARY KEY(y,r,c)
                )
            """)

            # Индексы для быстрой фильтрации по годам
            cur.execute("CREATE INDEX IF NOT EXISTS idx_employees_y ON employees(y)")
            cur.execute("CREATE INDEX IF NOT EXISTS idx_inventory_y ON inventory(y)")
            cur.execute("CREATE INDEX IF NOT EXISTS idx_ts_norms_y ON ts_norms_data(y)")


def load_state(year: int) -> dict:
    """Загрузка полного состояния справочника (нормы, сотрудники, локомотивы) за выбранный год."""
    norms = [[MONTHS[r], "", "", "", "", "", "", ""] for r in range(12)]
    employees = []
    inventory = []

    with DB_LOCK, conn_ctx() as conn:
        cur = conn.cursor()

        # 1. Загрузка норм времени
        for row in cur.execute("SELECT r, c, v FROM ts_norms_data WHERE y=?", (year,)):
            r, c = int(row["r"]), int(row["c"])
            if 0 <= r < 12 and 0 <= c < 8:
                norms[r][c] = text(row["v"])

        # 2. Загрузка сотрудников
        emp_idx = 0
        for row in cur.execute(
            """
            SELECT pos, name, full_name, tab_num, milk, milk_issue,
                   hire_date, exclude_date, vacation_days, milk_note
            FROM employees
            WHERE y=?
            ORDER BY rowid
            """,
            (year,),
        ):
            emp_idx += 1
            v_days = int(row["vacation_days"]) if row["vacation_days"] is not None else 28

            raw_tab = text(row["tab_num"])
            num_part = emp_idx
            if raw_tab and raw_tab.startswith("ID_"):
                digits = "".join(ch for ch in raw_tab if ch.isdigit())
                if digits:
                    num_part = int(digits)

            anon_id = f"ID_{num_part:03d}"
            anon_short = f"Работник №{num_part}"
            anon_full = f"СотрудникПолн №{num_part}"
            anon_pos = f"Должность №{num_part}"

            employees.append([
                anon_pos,
                anon_short,
                anon_full,
                anon_id,
                int(row["milk"] or 0),
                int(row["milk_issue"] or 0),
                text(row["hire_date"]),
                text(row["exclude_date"]),
                v_days,
                text(row["milk_note"])
            ])

        # 3. Загрузка локомотивов инвентарного парка
        inventory_rows = cur.execute(
            """
            SELECT ser, num, inv, COALESCE(sort_order, 0) AS sort_order,
                   COALESCE(updated_at, 0) AS updated_at, COALESCE(wheel_pair_count, 0) AS wheel_pair_count,
                   COALESCE(section_count, 0) AS section_count, COALESCE(deleted_at, 0) AS deleted_at,
                   COALESCE(eight_digit_number, '') AS eight_digit_number,
                   COALESCE(manufacture_year, '') AS manufacture_year,
                   COALESCE(service_life, '') AS service_life, rowid
            FROM inventory
            WHERE y=?
            ORDER BY COALESCE(sort_order, 0) ASC, COALESCE(updated_at, 0) DESC, rowid
            """,
            (year,),
        ).fetchall()

        best_inventory_rows: dict[str, sqlite3.Row] = {}
        for row in inventory_rows:
            ser = text(row["ser"]).strip()
            num = text(row["num"]).strip()
            if not (ser or num):
                continue
            key = f"{ser.upper()}|{num}"
            current = best_inventory_rows.get(key)
            if current is None or (
                int(row["updated_at"] or 0),
                int(row["deleted_at"] or 0),
                int(row["sort_order"] or 0),
                int(row["rowid"] or 0),
            ) > (
                int(current["updated_at"] or 0),
                int(current["deleted_at"] or 0),
                int(current["sort_order"] or 0),
                int(current["rowid"] or 0),
            ):
                best_inventory_rows[key] = row

        for row in sorted(
            best_inventory_rows.values(),
            key=lambda item: (
                int(item["sort_order"] or 0),
                -int(item["updated_at"] or 0),
                -int(item["deleted_at"] or 0),
                text(item["ser"]),
                text(item["num"]),
            ),
        ):
            inventory.append([
                text(row["ser"]),
                text(row["num"]),
                text(row["inv"]),
                int(row["wheel_pair_count"] or 0),
                int(row["section_count"] or 0),
                text(row["eight_digit_number"]),
                text(row["manufacture_year"]),
                text(row["service_life"]),
                int(row["deleted_at"] or 0),
            ])

    return {"year": year, "norms": norms, "employees": employees, "inventory": inventory}


def lookup_private_employee(id_input: str) -> dict:
    """Поиск данных сотрудника по табельному номеру, ID или ФИО в локальном файле."""
    raw = text(id_input).strip()
    if not raw:
        return {}

    raw_lower = raw.lower()
    raw_digits = "".join(c for c in raw if c.isdigit())
    raw_int = int(raw_digits) if raw_digits else None

    csv_file = ROOT.parent / "employees_private.csv"
    xlsx_file = ROOT.parent / "employees_private.xlsx"

    rows = []
    if csv_file.exists():
        try:
            import csv
            with open(csv_file, "r", encoding="utf-8-sig") as f:
                for r in csv.reader(f, delimiter=";"):
                    if r and not r[0].startswith("Код"):
                        rows.append([text(x).strip() for x in r])
        except Exception:
            pass
    elif xlsx_file.exists():
        try:
            import openpyxl
            wb = openpyxl.load_workbook(xlsx_file)
            ws = wb.active
            for r in ws.iter_rows(values_only=True):
                if r and r[0] and not text(r[0]).startswith("Код"):
                    rows.append([text(x).strip() for x in r])
        except Exception:
            pass

    for r in rows:
        r_id = r[0] if len(r) > 0 else ""
        r_pos = r[1] if len(r) > 1 else ""
        r_full = r[2] if len(r) > 2 else ""
        r_short = r[3] if len(r) > 3 else r_full
        r_tab = r[4] if len(r) > 4 else ""

        r_num_digits = "".join(c for c in r_id if c.isdigit())
        r_int = int(r_num_digits) if r_num_digits else None
        r_tab_digits = "".join(c for c in r_tab if c.isdigit())
        r_tab_int = int(r_tab_digits) if r_tab_digits else None

        # 1. Совпадение по анонимному ID (ID_001)
        if raw.upper() == r_id.upper():
            return {"pos": r_pos, "name": r_short, "full_name": r_full, "tab_num": r_id, "found": True}

        # 2. Совпадение по реальному табельному номеру
        if r_tab and (raw == r_tab or (raw_int is not None and r_tab_int is not None and raw_int == r_tab_int)):
            return {"pos": r_pos, "name": r_short, "full_name": r_full, "tab_num": r_id, "found": True}

        # 3. Совпадение по порядковому номеру (например 1, 001)
        if raw_int is not None and r_int is not None and raw_int == r_int:
            return {"pos": r_pos, "name": r_short, "full_name": r_full, "tab_num": r_id, "found": True}

        # 4. Совпадение по подстроке ФИО
        if raw_lower and (raw_lower in r_short.lower() or raw_lower in r_full.lower()):
            return {"pos": r_pos, "name": r_short, "full_name": r_full, "tab_num": r_id, "found": True}

    # Если в файле нет, но передан номер < 1000 — возвращаем канонический ID_00X
    if raw_int is not None and raw_int < 1000:
        canon_id = f"ID_{raw_int:03d}"
        return {
            "pos": "",
            "name": "",
            "full_name": "",
            "tab_num": canon_id,
            "found": False,
        }

    return {"tab_num": raw, "name": raw, "full_name": raw, "pos": "", "found": False}


def save_state(payload: dict) -> None:
    """Атомарное сохранение изменений справочника за указанный год."""
    year = int(payload.get("year") or dt.date.today().year)
    norms = payload.get("norms") or []
    employees = payload.get("employees") or []
    inventory = payload.get("inventory") or []

    with DB_LOCK, conn_ctx() as conn:
        with conn:
            cur = conn.cursor()

            # 1. Сохранение норм времени
            cur.execute("DELETE FROM ts_norms_data WHERE y=?", (year,))
            norm_rows = []
            for r, row in enumerate(norms[:12]):
                for c, value in enumerate((row or [])[:8]):
                    value = text(value).strip()
                    if value:
                        norm_rows.append((year, r, c, value))
            cur.executemany("INSERT INTO ts_norms_data VALUES (?,?,?,?)", norm_rows)

            # 2. Сохранение сотрудников
            cur.execute("DELETE FROM employees WHERE y=?", (year,))
            emp_rows = []
            for row in employees:
                row = list(row or []) + [""] * 10
                pos, name, full_name, tab_num = [text(v).strip() for v in row[:4]]

                # Автоподтягивание данных сотрудника по табельному/ID при необходимости
                search_key = tab_num or name or full_name
                if search_key:
                    info = lookup_private_employee(search_key)
                    if info:
                        if not pos or pos.startswith("Должность №"):
                            pos = info.get("pos") or pos
                        if not name or name.startswith("Работник №"):
                            name = info.get("name") or name
                        if not full_name or full_name.startswith("СотрудникПолн №"):
                            full_name = info.get("full_name") or full_name
                        if not tab_num or tab_num.isdigit():
                            tab_num = info.get("tab_num") or tab_num

                milk = 1 if row[4] else 0
                milk_issue = 1 if row[5] else 0
                hire_date = text(row[6]).strip()
                exclude_date = text(row[7]).strip()
                try:
                    v_str = text(row[8]).strip()
                    vacation_days = int(v_str) if v_str != "" else 28
                except Exception:
                    vacation_days = 28
                milk_note = text(row[9]).strip()

                if pos or name or tab_num:
                    emp_rows.append((
                        year, pos, name, full_name, tab_num, milk, milk_issue,
                        hire_date, exclude_date, vacation_days, milk_note
                    ))

            cur.executemany(
                """
                INSERT INTO employees (y, pos, name, full_name, tab_num, milk, milk_issue, hire_date, exclude_date, vacation_days, milk_note)
                VALUES (?,?,?,?,?,?,?,?,?,?,?)
                """,
                emp_rows,
            )

            # 3. Обновление инвентарного парка локомотивов
            existing_rows = cur.execute(
                """
                SELECT ser, num, inv, COALESCE(sort_order, 0) AS sort_order,
                       COALESCE(updated_at, 0) AS updated_at, COALESCE(deleted_at, 0) AS deleted_at,
                       COALESCE(wheel_pair_count, 0) AS wheel_pair_count, COALESCE(section_count, 0) AS section_count,
                       COALESCE(eight_digit_number, '') AS eight_digit_number,
                       COALESCE(manufacture_year, '') AS manufacture_year,
                       COALESCE(service_life, '') AS service_life, rowid
                FROM inventory
                WHERE y=?
                """,
                (year,),
            ).fetchall()

            existing_map: dict[tuple[str, str], sqlite3.Row] = {}
            duplicate_rowids: list[int] = []
            for row in sorted(
                existing_rows,
                key=lambda item: (
                    -int(item["updated_at"] or 0),
                    -int(item["deleted_at"] or 0),
                    int(item["sort_order"] or 0),
                    int(item["rowid"] or 0),
                ),
            ):
                ser = text(row["ser"]).strip()
                num = text(row["num"]).strip()
                if not (ser or num):
                    continue
                key = (ser.upper(), num)
                if key in existing_map:
                    duplicate_rowids.append(int(row["rowid"]))
                    continue
                existing_map[key] = row

            if duplicate_rowids:
                placeholders = ",".join("?" for _ in duplicate_rowids)
                cur.execute(f"DELETE FROM inventory WHERE rowid IN ({placeholders})", duplicate_rowids)

            submitted_keys: set[tuple[str, str]] = set()
            now = int(dt.datetime.now().timestamp() * 1000)

            for order_index, row in enumerate(inventory, start=1):
                raw_row = list(row or [])
                legacy_format = len(raw_row) < 9
                row = raw_row + [""] * max(0, 9 - len(raw_row))
                ser, num = [text(v).strip() for v in row[:2]]
                try:
                    wheel_pair_count = int(row[3] or 0)
                except Exception:
                    wheel_pair_count = 0
                try:
                    section_count = int(row[4] or 0)
                except Exception:
                    section_count = 0
                inv = text(row[2]).strip()
                eight_digit_number = text(row[5]).strip()
                manufacture_year = "" if legacy_format else text(row[6]).strip()
                service_life = "" if legacy_format else text(row[7]).strip()
                try:
                    deleted_at = int((row[7] if legacy_format else row[8]) or 0)
                except Exception:
                    deleted_at = 0

                if not (ser or num):
                    continue
                key = (ser.upper(), num)
                submitted_keys.add(key)

                if key in existing_map:
                    existing_inv = text(existing_map[key]["inv"]).strip()
                    inv_value = inv if inv else existing_inv
                    cur.execute(
                        """
                        UPDATE inventory
                        SET inv=?, wheel_pair_count=?, section_count=?, eight_digit_number=?,
                            manufacture_year=?, service_life=?, sort_order=?, updated_at=?, deleted_at=?
                        WHERE y=? AND ser=? AND num=?
                        """,
                        (
                            inv_value,
                            wheel_pair_count or None,
                            section_count or None,
                            eight_digit_number,
                            manufacture_year,
                            service_life,
                            order_index,
                            now,
                            deleted_at,
                            year,
                            text(existing_map[key]["ser"]).strip(),
                            num,
                        ),
                    )
                else:
                    cur.execute(
                        """
                        INSERT INTO inventory (
                            y, ser, num, inv, wheel_pair_count, section_count,
                            eight_digit_number, manufacture_year, service_life,
                            sort_order, updated_at, deleted_at
                        )
                        VALUES (?,?,?,?,?,?,?,?,?,?,?,?)
                        """,
                        (
                            year, ser.upper(), num, inv,
                            wheel_pair_count or None, section_count or None,
                            eight_digit_number, manufacture_year, service_life,
                            order_index, now, deleted_at
                        ),
                    )

            # Помечаем удаленными те строки, которые были исключены пользователем
            for (ser_key, num_key), row in existing_map.items():
                if (ser_key, num_key) in submitted_keys:
                    continue
                if int(row["deleted_at"] or 0) > 0:
                    continue
                cur.execute(
                    """
                    UPDATE inventory
                    SET updated_at=?, deleted_at=?, sort_order=COALESCE(sort_order, ?)
                    WHERE y=? AND ser=? AND num=?
                    """,
                    (now, now, len(inventory) + 1, year, text(row["ser"]).strip(), text(row["num"]).strip()),
                )


def purge_deleted_inventory() -> dict:
    """Полное удаление из базы помеченных на удаление локомотивов и связанных замеров КП."""
    with DB_LOCK, conn_ctx() as conn:
        with conn:
            cur = conn.cursor()
            deleted_rows = cur.execute(
                """
                SELECT y, ser, num
                FROM inventory
                WHERE COALESCE(deleted_at, 0) > 0
                """
            ).fetchall()
            if not deleted_rows:
                return {"ok": True, "purged": 0}

            numbers = [text(row["num"]).strip() for row in deleted_rows if text(row["num"]).strip()]
            cur.execute("DELETE FROM inventory WHERE COALESCE(deleted_at, 0) > 0")
            if numbers:
                placeholders = ",".join("?" for _ in numbers)
                cur.execute(
                    f"DELETE FROM kp_data WHERE TRIM(COALESCE(locomotive, '')) IN ({placeholders})",
                    numbers
                )
            return {"ok": True, "purged": len(deleted_rows)}


def purge_inventory_row(year: int, ser: str, num: str) -> dict:
    """Удаление одного конкретного локомотива по году, серии и номеру."""
    year = int(year)
    ser = text(ser).strip()
    num = text(num).strip()
    if not ser or not num:
        return {"ok": False, "error": "Некорректные данные локомотива."}

    with DB_LOCK, conn_ctx() as conn:
        with conn:
            cur = conn.cursor()
            row = cur.execute(
                """
                SELECT y, ser, num
                FROM inventory
                WHERE y=? AND TRIM(COALESCE(ser, ''))=? AND TRIM(COALESCE(num, ''))=?
                LIMIT 1
                """,
                (year, ser, num),
            ).fetchone()
            if not row:
                return {"ok": True, "purged": 0}

            cur.execute(
                "DELETE FROM inventory WHERE y=? AND TRIM(COALESCE(ser, ''))=? AND TRIM(COALESCE(num, ''))=?",
                (year, ser, num),
            )
            cur.execute("DELETE FROM kp_data WHERE TRIM(COALESCE(locomotive, ''))=?", (num,))
            return {"ok": True, "purged": 1}
