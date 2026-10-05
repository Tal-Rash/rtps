from __future__ import annotations

import datetime as dt
import io
from http import HTTPStatus
from pathlib import Path

# Константы и методы работы с базой данных
from constants import (
    ARCHIVE_EXCEL_HEADERS,
    DB_LOCK,
    DEFAULT_REPAIR_OPTIONS,
    WEAR_TREND_METRICS,
)
from storage import (
    connect,
    default_section_count,
    ensure_import_locomotive,
    load_archive_rows,
    load_effective_kp_values_for_date,
    locomotive_axis_count,
    normalize_repair_type,
    normalize_text,
    parse_excel_int,
    parse_float_value,
    resolve_archive_diameter_pair,
    series_for_locomotive,
    text,
)


def require_openpyxl():
    """Ленивый импорт openpyxl для экономии памяти при запуске."""
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError as exc:
        raise RuntimeError("Для работы с Excel нужен пакет openpyxl.") from exc
    return Workbook, load_workbook, Alignment, Border, Font, PatternFill, Side, DataValidation


def build_archive_workbook():
    """Создание книги Excel со стилизованной шапкой для архива замеров."""
    Workbook, _, Alignment, Border, Font, PatternFill, Side, _ = require_openpyxl()
    wb = Workbook()
    ws = wb.active
    ws.title = "Архив"
    ws.append(ARCHIVE_EXCEL_HEADERS)

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    thin = Side(style="thin", color="BFBFBF")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

    widths = [14, 12, 14, 12, 10, 10, 14, 14, 18, 18, 18, 18, 18, 18, 18, 18]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(1, index).column_letter].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:P1"
    ws.row_dimensions[1].height = 34

    return wb


def archive_excel_template_bytes() -> bytes:
    """Генерация пустого шаблона Excel для загрузки архива."""
    wb = build_archive_workbook()
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def parse_excel_date(value) -> str:
    """Нормализация даты из ячейки Excel в ISO формат YYYY-MM-DD."""
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    raw = text(value).strip()
    if not raw:
        return ""
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d.%m.%y", "%Y/%m/%d"):
        try:
            return dt.datetime.strptime(raw, fmt).date().isoformat()
        except ValueError:
            pass
    return raw


def format_excel_export_date(value: str) -> str:
    """Форматирование даты в формат ДД.ММ.ГГГГ для экспорта в Excel."""
    raw = text(value).strip()
    if not raw:
        return ""
    try:
        return dt.datetime.strptime(raw, "%Y-%m-%d").strftime("%d.%m.%Y")
    except ValueError:
        return raw


def excel_cell_text(value) -> str:
    """Приведение значения ячейки Excel к строковому типу."""
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    return text(value).strip()


def excel_num_text(value) -> str:
    """Форматирование чисел с запятой для корректного отображения в русской локали Excel."""
    raw = text(value).strip()
    return raw.replace(".", ",") if raw else ""


def normalize_excel_header(value) -> str:
    """Очистка заголовка колонки Excel от пробелов и спецсимволов для сопоставления."""
    return "".join(ch for ch in text(value).strip().lower() if ch.isalnum())


def build_archive_export_rows(selected_locomotives: list[str] | None = None, date_from: str = "", date_to: str = "") -> list[list[str]]:
    """Формирование плоских строк архива замеров для записи в Excel."""
    locomotives_filter = {text(item).strip() for item in selected_locomotives or [] if text(item).strip()}
    date_from = text(date_from).strip()
    date_to = text(date_to).strip()
    query = """
        SELECT y, measurement_date, locomotive, repair_type, r, c, v
        FROM archive_data
        WHERE TRIM(COALESCE(measurement_date, '')) <> ''
    """
    params: list[str] = []
    if locomotives_filter:
        query += f" AND locomotive IN ({','.join('?' for _ in locomotives_filter)})"
        params.extend(sorted(locomotives_filter))
    if date_from:
        query += " AND measurement_date >= ?"
        params.append(date_from)
    if date_to:
        query += " AND measurement_date <= ?"
        params.append(date_to)
    query += " ORDER BY y, measurement_date, locomotive, repair_type, r, c"

    with DB_LOCK, connect() as conn:
        cur = conn.cursor()
        rows = cur.execute(query, params).fetchall()
        grouped: dict[tuple[int, str, str, str, int], list[str]] = {}
        for row in rows:
            r = int(row["r"])
            if r < 2:
                continue
            key = (int(row["y"] or 0), text(row["measurement_date"]), text(row["locomotive"]), normalize_repair_type(row["repair_type"]), r)
            grouped.setdefault(key, [""] * 12)
            c = int(row["c"])
            if 0 <= c < 12:
                grouped[key][c] = text(row["v"])

        kp_cache: dict[tuple[str, str], dict[tuple[int, int], str]] = {}
        export_rows: list[list[str]] = []
        for (_, measurement_date, locomotive, repair_type, r), values in sorted(grouped.items()):
            series = series_for_locomotive(cur, locomotive)
            bandage_left, bandage_right, diameter_left, diameter_right = resolve_archive_diameter_pair(
                cur,
                locomotive,
                measurement_date,
                r - 2,
                values,
                kp_cache,
            )

            export_rows.append(
                [
                    format_excel_export_date(measurement_date),
                    locomotive,
                    normalize_repair_type(repair_type),
                    series,
                    values[0] or "1",
                    values[1] or str(r - 1),
                    excel_num_text(values[2]),
                    excel_num_text(values[3]),
                    excel_num_text(values[4]),
                    excel_num_text(values[5]),
                    excel_num_text(values[6]),
                    excel_num_text(values[7]),
                    excel_num_text(bandage_left),
                    excel_num_text(bandage_right),
                    excel_num_text(diameter_left),
                    excel_num_text(diameter_right),
                ]
            )
        return export_rows


def archive_excel_export_bytes(selected_locomotives: list[str] | None = None, date_from: str = "", date_to: str = "") -> tuple[bytes, int]:
    """Сборка готового байтового файла Excel с архивом замеров."""
    _, _, Alignment, _, _, _, _, _ = require_openpyxl()

    wb = build_archive_workbook()
    ws = wb.active
    rows = build_archive_export_rows(selected_locomotives, date_from, date_to)
    out_row = 2
    for row_data in rows:
        for col, value in enumerate(row_data, start=1):
            cell = ws.cell(out_row, col, value)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        out_row += 1
    ws.auto_filter.ref = f"A1:P{max(1, out_row - 1)}"
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue(), len(rows)


def import_archive_excel_bytes(data: bytes) -> dict:
    """Импорт и валидация замеров колёсных пар из входящего Excel-файла."""
    _, load_workbook, *_ = require_openpyxl()
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active

    aliases = {
        "measurement_date": {"датазамера", "датавыполненияобмера", "датаобмера", "measurementdate"},
        "series": {"серия", "series"},
        "locomotive": {"локомотив", "номерлокомотива", "locomotive"},
        "repair_type": {"видремонта", "видрем", "repairtype"},
        "wheel_pair": {"номеркп", "wheelpair", "wheelpairnumber", "ось", "номероси", "kp"},
        "section": {"секция", "вагон", "section"},
        "prokat_left": {"прокатлев", "левпрокат", "prokatleft", "flangewearleft"},
        "prokat_right": {"прокатправ", "правпрокат", "prokatright", "flangewearright"},
        "greben_left": {"толщинагребнялев", "левтолщинагребня", "grebenleft", "flangethicknessleft"},
        "greben_right": {"толщинагребняправ", "правтолщинагребня", "grebenright", "flangethicknessright"},
        "krut_left": {"крутизнагребнялев", "левкрутизнагребня", "krutleft", "flangesteepnessleft"},
        "krut_right": {"крутизнагребняправ", "правкрутизнагребня", "krutright", "flangesteepnessright"},
        "bandage_thickness_left": {"толщинабандажалева", "толщинабандажалев", "leftbandagethickness", "bandagethicknessleft"},
        "bandage_thickness_right": {"толщинабандажаправ", "rightbandagethickness", "bandagethicknessright"},
        "bandage_diameter_left": {"диаметрбандажалева", "диаметрбандажалев", "leftbandagediameter", "bandagediameterleft"},
        "bandage_diameter_right": {"диаметрбандажаправ", "rightbandagediameter", "bandagediameterright"},
    }

    header_row_index = None
    header_columns: dict[str, int] = {}
    for row_index, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
        normalized = [normalize_excel_header(value) for value in row]
        found: dict[str, int] = {}
        for key, alias_set in aliases.items():
            for idx, header in enumerate(normalized):
                if header in alias_set:
                    found[key] = idx
                    break
        if len(found) >= 5:
            header_row_index = row_index
            header_columns = found
            break

    if header_row_index is None:
        return {"error": "Не найдена строка заголовков."}, HTTPStatus.BAD_REQUEST

    missing_headers = [name for name in ("measurement_date", "locomotive", "repair_type", "wheel_pair") if name not in header_columns]
    if missing_headers:
        return {"error": "В Excel не найдены обязательные колонки: " + ", ".join(missing_headers)}, HTTPStatus.BAD_REQUEST

    required_metric_columns = {
        "prokat_left": 2,
        "prokat_right": 3,
        "greben_left": 4,
        "greben_right": 5,
        "krut_left": 6,
        "krut_right": 7,
        "bandage_thickness_left": 8,
        "bandage_thickness_right": 9,
    }
    optional_metric_columns = {"bandage_diameter_left": 10, "bandage_diameter_right": 11}

    measurements: dict[tuple[str, str, str], dict] = {}
    errors: list[str] = []
    current_series = current_date = current_loco = current_repair = current_section = ""

    def row_value(row, column_name: str) -> str:
        idx = header_columns.get(column_name)
        if idx is None or idx >= len(row):
            return ""
        return excel_cell_text(row[idx])

    for row_index, row in enumerate(ws.iter_rows(min_row=header_row_index + 1, values_only=True), start=header_row_index + 1):
        if not row or all(cell is None or text(cell).strip() == "" for cell in row):
            continue
        row_date = parse_excel_date(row_value(row, "measurement_date")) or current_date
        row_loco = row_value(row, "locomotive") or current_loco
        row_repair = row_value(row, "repair_type") or current_repair
        row_series = row_value(row, "series") or current_series
        row_section = row_value(row, "section") or current_section
        wheel_pair_number = parse_excel_int(row_value(row, "wheel_pair"))

        if row_value(row, "measurement_date"):
            current_date = row_date
        if row_value(row, "locomotive"):
            current_loco = row_loco
        if row_value(row, "repair_type"):
            current_repair = row_repair
        if row_value(row, "series"):
            current_series = row_series
        if row_value(row, "section"):
            current_section = row_section

        if not row_date or not row_loco or not row_repair:
            errors.append(f"Строка {row_index}: не заполнены дата, локомотив или вид ремонта.")
            continue
        if wheel_pair_number is None or wheel_pair_number <= 0:
            errors.append(f"Строка {row_index}: не указан корректный номер КП.")
            continue

        values: dict[int, str] = {}
        missing_fields = []
        for column_name, db_col in required_metric_columns.items():
            value = row_value(row, column_name)
            if value == "":
                missing_fields.append(column_name)
            else:
                values[db_col] = value
        for column_name, db_col in optional_metric_columns.items():
            value = row_value(row, column_name)
            if value != "":
                values[db_col] = value
        if missing_fields:
            errors.append(f"Строка {row_index}: не заполнены поля " + ", ".join(missing_fields))
            continue

        group = measurements.setdefault((row_date, row_loco, row_repair), {"series": row_series, "rows": {}, "sections": {}})
        if row_series and not group["series"]:
            group["series"] = row_series
        if row_section:
            group["sections"][wheel_pair_number] = row_section
        group["rows"][wheel_pair_number] = values

    if not measurements:
        return {"error": "Не удалось импортировать ни одного замера.", "errors": errors[:20]}, HTTPStatus.BAD_REQUEST

    imported_measurements = 0
    imported_cells = 0
    with DB_LOCK, connect() as conn:
        cur = conn.cursor()
        kp_cache: dict[str, dict[tuple[int, int], str]] = {}
        cur.execute("BEGIN")
        for (measurement_date, locomotive, repair_type), meta in sorted(measurements.items()):
            wheel_pair_numbers = sorted(meta["rows"].keys())
            if not wheel_pair_numbers:
                continue
            year = int(measurement_date[:4]) if len(measurement_date) >= 4 and measurement_date[:4].isdigit() else dt.date.today().year
            series = meta["series"] or series_for_locomotive(cur, locomotive)
            ensure_import_locomotive(cur, series, locomotive, max(wheel_pair_numbers))
            if locomotive not in kp_cache:
                kp_rows = cur.execute("SELECT r, c, v FROM kp_data WHERE locomotive=? AND c IN (2, 3)", (locomotive,)).fetchall()
                kp_cache[locomotive] = {(int(row["r"]), int(row["c"])): text(row["v"]).strip() for row in kp_rows}

            axis_count = locomotive_axis_count(series_for_locomotive(cur, locomotive), locomotive)
            db_rows: list[tuple[int, str, str, str, int, int, str]] = []
            for wheel_pair_number in wheel_pair_numbers:
                row_values = meta["rows"][wheel_pair_number]
                table_row = wheel_pair_number + 1
                section_value = meta["sections"].get(wheel_pair_number, "1" if axis_count == 6 else str(((wheel_pair_number - 1) // 4) + 1))
                kp_index = wheel_pair_number - 1
                bandage_left = row_values.get(8, "")
                bandage_right = row_values.get(9, "")
                diameter_left = row_values.get(10, "")
                diameter_right = row_values.get(11, "")
                if not diameter_left:
                    kp_left = parse_float_value(kp_cache[locomotive].get((kp_index, 2), ""))
                    bandage_left_value = parse_float_value(bandage_left)
                    if kp_left is not None and bandage_left_value is not None:
                        diameter_left = str(int(round(kp_left + bandage_left_value * 2)))
                if not diameter_right:
                    kp_right = parse_float_value(kp_cache[locomotive].get((kp_index, 3), ""))
                    bandage_right_value = parse_float_value(bandage_right)
                    if kp_right is not None and bandage_right_value is not None:
                        diameter_right = str(int(round(kp_right + bandage_right_value * 2)))

                full_values = {
                    0: section_value,
                    1: str(wheel_pair_number),
                    2: row_values.get(2, ""),
                    3: row_values.get(3, ""),
                    4: row_values.get(4, ""),
                    5: row_values.get(5, ""),
                    6: row_values.get(6, ""),
                    7: row_values.get(7, ""),
                    8: bandage_left,
                    9: bandage_right,
                    10: diameter_left,
                    11: diameter_right,
                }
                for col, value in full_values.items():
                    if value != "":
                        db_rows.append((year, measurement_date, locomotive, repair_type, table_row, col, value))
            if not db_rows:
                continue
            cur.execute(
                "DELETE FROM archive_data WHERE y=? AND measurement_date=? AND locomotive=? AND repair_type=?",
                (year, measurement_date, locomotive, repair_type),
            )
            cur.executemany(
                "INSERT OR REPLACE INTO archive_data (y, measurement_date, locomotive, repair_type, r, c, v) VALUES (?, ?, ?, ?, ?, ?, ?)",
                db_rows,
            )
            imported_measurements += 1
            imported_cells += len(db_rows)
        conn.commit()

    return {"ok": True, "imported_measurements": imported_measurements, "imported_cells": imported_cells, "errors": errors[:10]}
