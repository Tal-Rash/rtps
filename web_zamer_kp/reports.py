from __future__ import annotations

import base64
import datetime as dt
import io
import json
import re
import sqlite3
import zlib
from http import HTTPStatus
from pathlib import Path

from constants import (
    ROOT,
    DB_LOCK,
    ARCHIVE_EXCEL_HEADERS,
    WEAR_TREND_METRICS,
    DEFAULT_REPAIR_OPTIONS,
)
from storage import (
    connect,
    text,
    normalize_text,
    normalize_repair_type,
    parse_excel_int,
    parse_float_value,
    series_for_locomotive,
    locomotive_axis_count,
    default_section_count,
    load_inventory_records,
    upsert_inventory_locomotive,
    ensure_import_locomotive,
    load_effective_kp_values_for_date,
    resolve_archive_diameter_pair,
    load_archive_rows,
)

def require_openpyxl():
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError as exc:
        raise RuntimeError("Для работы с Excel нужен пакет openpyxl.") from exc
    return Workbook, load_workbook, Alignment, Border, Font, PatternFill, Side, DataValidation


def build_archive_workbook():
    Workbook, _, Alignment, Border, Font, PatternFill, Side, DataValidation = require_openpyxl()
    wb = Workbook()
    ws = wb.active
    ws.title = "Архив"
    ws.append(ARCHIVE_EXCEL_HEADERS)

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    thin = Side(style="thin", color="BFBFBF")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

    widths = [14, 12, 14, 12, 10, 10, 14, 14, 18, 18, 18, 18, 18, 18, 18, 18]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(1, index).column_letter].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:P1"
    ws.row_dimensions[1].height = 34

    return wb


def archive_excel_template_bytes() -> bytes:
    wb = build_archive_workbook()
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def parse_excel_date(value) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    raw = text(value).strip()
    if not raw:
        return ""
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d.%m.%y", "%Y/%m/%d"):
        try:
            return dt.datetime.strptime(raw, fmt).date().isoformat()
        except ValueError:
            pass
    return raw


def format_excel_export_date(value: str) -> str:
    raw = text(value).strip()
    if not raw:
        return ""
    try:
        return dt.datetime.strptime(raw, "%Y-%m-%d").strftime("%d.%m.%Y")
    except ValueError:
        return raw


def excel_cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    return text(value).strip()


def excel_num_text(value) -> str:
    raw = text(value).strip()
    return raw.replace(".", ",") if raw else ""


def normalize_excel_header(value) -> str:
    return "".join(ch for ch in text(value).strip().lower() if ch.isalnum())


def build_archive_export_rows(selected_locomotives: list[str] | None = None, date_from: str = "", date_to: str = "") -> list[list[str]]:
    locomotives_filter = {text(item).strip() for item in selected_locomotives or [] if text(item).strip()}
    date_from = text(date_from).strip()
    date_to = text(date_to).strip()
    query = """
        SELECT y, measurement_date, locomotive, repair_type, r, c, v
        FROM archive_data
        WHERE TRIM(COALESCE(measurement_date, '')) <> ''
    """
    params: list[str] = []
    if locomotives_filter:
        query += f" AND locomotive IN ({','.join('?' for _ in locomotives_filter)})"
        params.extend(sorted(locomotives_filter))
    if date_from:
        query += " AND measurement_date >= ?"
        params.append(date_from)
    if date_to:
        query += " AND measurement_date <= ?"
        params.append(date_to)
    query += " ORDER BY y, measurement_date, locomotive, repair_type, r, c"

    with DB_LOCK, connect() as conn:
        cur = conn.cursor()
        rows = cur.execute(query, params).fetchall()
        grouped: dict[tuple[int, str, str, str, int], list[str]] = {}
        for row in rows:
            r = int(row["r"])
            if r < 2:
                continue
            key = (int(row["y"] or 0), text(row["measurement_date"]), text(row["locomotive"]), normalize_repair_type(row["repair_type"]), r)
            grouped.setdefault(key, [""] * 12)
            c = int(row["c"])
            if 0 <= c < 12:
                grouped[key][c] = text(row["v"])

        kp_cache: dict[tuple[str, str], dict[tuple[int, int], str]] = {}
        export_rows: list[list[str]] = []
        for (_, measurement_date, locomotive, repair_type, r), values in sorted(grouped.items()):
            series = series_for_locomotive(cur, locomotive)
            bandage_left, bandage_right, diameter_left, diameter_right = resolve_archive_diameter_pair(
                cur,
                locomotive,
                measurement_date,
                r - 2,
                values,
                kp_cache,
            )

            export_rows.append(
                [
                    format_excel_export_date(measurement_date),
                    locomotive,
                    normalize_repair_type(repair_type),
                    series,
                    values[0] or "1",
                    values[1] or str(r - 1),
                    excel_num_text(values[2]),
                    excel_num_text(values[3]),
                    excel_num_text(values[4]),
                    excel_num_text(values[5]),
                    excel_num_text(values[6]),
                    excel_num_text(values[7]),
                    excel_num_text(bandage_left),
                    excel_num_text(bandage_right),
                    excel_num_text(diameter_left),
                    excel_num_text(diameter_right),
                ]
            )
        return export_rows


def archive_excel_export_bytes(selected_locomotives: list[str] | None = None, date_from: str = "", date_to: str = "") -> tuple[bytes, int]:
    _, _, Alignment, _, _, _, _, _ = require_openpyxl()

    wb = build_archive_workbook()
    ws = wb.active
    rows = build_archive_export_rows(selected_locomotives, date_from, date_to)
    out_row = 2
    for row_data in rows:
        for col, value in enumerate(row_data, start=1):
            cell = ws.cell(out_row, col, value)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        out_row += 1
    ws.auto_filter.ref = f"A1:P{max(1, out_row - 1)}"
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue(), len(rows)


def phone_reference_export_payload() -> dict:
    with DB_LOCK, connect() as conn:
        cur = conn.cursor()
        locomotives = load_inventory_records(cur, include_deleted=True)
        kp_rows = cur.execute(
            "SELECT locomotive, r, c, v FROM kp_data WHERE TRIM(COALESCE(locomotive, '')) <> '' ORDER BY locomotive, r, c"
        ).fetchall()

    kp_map: dict[str, dict[int, dict[int, str]]] = {}
    for row in kp_rows:
        locomotive = text(row["locomotive"]).strip()
        kp_map.setdefault(locomotive, {}).setdefault(int(row["r"]), {})[int(row["c"])] = text(row["v"])

    payload_locomotives: list[dict[str, object]] = []
    for locomotive in locomotives:
        number = text(locomotive.get("number")).strip()
        series = text(locomotive.get("series")).strip()
        wheel_pair_count = max(1, locomotive_axis_count(series, number))
        row_map = kp_map.get(number, {})
        wheel_pairs: list[dict[str, object]] = []
        for pair_index in range(wheel_pair_count):
            row = row_map.get(pair_index, {})
            wheel_pairs.append(
                {
                    "number": pair_index + 1,
                    "axisNumber": parse_excel_int(row.get(1, "")) or (pair_index + 1),
                    "diameterLeft": parse_float_value(row.get(2, "")),
                    "diameterRight": parse_float_value(row.get(3, "")),
                }
            )
        payload_locomotives.append(
            {
                "series": series,
                "number": number,
                "wheelPairCount": wheel_pair_count,
                "inventoryNumber": text(locomotive.get("inventoryNumber") or locomotive.get("inv") or ""),
                "invNumber": text(locomotive.get("inventoryNumber") or locomotive.get("inv") or ""),
                "eightDigitNumber": text(locomotive.get("eightDigitNumber") or ""),
                "manufactureYear": text(locomotive.get("manufactureYear") or locomotive.get("manufacture_year") or ""),
                "sortOrder": int(locomotive.get("sortOrder") or 0),
                "updatedAt": int(locomotive.get("updatedAt") or 0),
                "deletedAt": int(locomotive.get("deletedAt") or 0),
                "wheelPairs": wheel_pairs,
            }
        )

    return {
        "formatVersion": 2,
        "exportType": "referenceData",
        "exportedAt": dt.datetime.now().isoformat(timespec="seconds"),
        "locomotives": payload_locomotives,
    }


def phone_archive_export_payload(selected_locomotives: list[str] | None = None, date_from: str = "", date_to: str = "") -> dict:
    locomotives_filter = {text(item).strip() for item in selected_locomotives or [] if text(item).strip()}
    date_from = text(date_from).strip()
    date_to = text(date_to).strip()
    query = """
        SELECT y, measurement_date, locomotive, repair_type, r, c, v
        FROM archive_data
        WHERE TRIM(COALESCE(measurement_date, '')) <> ''
    """
    params: list[str] = []
    if locomotives_filter:
        query += f" AND locomotive IN ({','.join('?' for _ in locomotives_filter)})"
        params.extend(sorted(locomotives_filter))
    if date_from:
        query += " AND measurement_date >= ?"
        params.append(date_from)
    if date_to:
        query += " AND measurement_date <= ?"
        params.append(date_to)
    query += " ORDER BY y, measurement_date, locomotive, repair_type, r, c"

    with DB_LOCK, connect() as conn:
        cur = conn.cursor()
        rows = cur.execute(query, params).fetchall()
        grouped: dict[tuple[int, str, str, str], dict[int, list[str]]] = {}
        series_cache: dict[str, str] = {}
        kp_cache: dict[tuple[str, str], dict[tuple[int, int], str]] = {}
        for row in rows:
            key = (
                int(row["y"] or 0),
                text(row["measurement_date"]).strip(),
                text(row["locomotive"]).strip(),
                normalize_repair_type(row["repair_type"]),
            )
            grouped.setdefault(key, {})[int(row["r"])] = grouped.setdefault(key, {}).get(int(row["r"]), [""] * 12)
            grouped[key][int(row["r"])][int(row["c"])] = text(row["v"])

        archive_items: list[dict[str, object]] = []
        for (year, measurement_date, locomotive, repair_type), rows_by_r in sorted(grouped.items()):
            series = series_cache.get(locomotive)
            if series is None:
                series = series_for_locomotive(cur, locomotive)
                series_cache[locomotive] = series
            wheel_pair_count = max(1, locomotive_axis_count(series, locomotive))
            wheel_pairs: list[dict[str, object]] = []
            for row_index in sorted(rows_by_r):
                row = rows_by_r[row_index]
                pair_number = parse_excel_int(row[1]) or max(1, row_index - 1)
                left_band, right_band, diameter_left, diameter_right = resolve_archive_diameter_pair(
                    cur,
                    locomotive,
                    measurement_date,
                    pair_number - 1,
                    row,
                    kp_cache,
                )

                wheel_pairs.append(
                    {
                        "number": pair_number,
                        "left": {
                            "flangeThickness": parse_float_value(row[4]),
                            "flangeWear": parse_float_value(row[2]),
                            "flangeSteepness": parse_float_value(row[6]),
                            "bandageThickness": parse_float_value(left_band),
                            "bandageDiameter": parse_float_value(diameter_left),
                        },
                        "right": {
                            "flangeThickness": parse_float_value(row[5]),
                            "flangeWear": parse_float_value(row[3]),
                            "flangeSteepness": parse_float_value(row[7]),
                            "bandageThickness": parse_float_value(right_band),
                            "bandageDiameter": parse_float_value(diameter_right),
                        },
                    }
                )

            archive_items.append(
                {
                    "formatVersion": 1,
                    "createdAt": dt.datetime.now().isoformat(timespec="seconds"),
                    "measurementId": f"{year}:{measurement_date}:{locomotive}:{repair_type}",
                    "locomotive": {
                        "series": series,
                        "number": locomotive,
                        "wheelPairCount": wheel_pair_count,
                        "comment": "",
                        "isNew": False,
                    },
                    "repairType": repair_type,
                    "measurementDate": measurement_date,
                    "wheelPairs": wheel_pairs,
                }
            )

    return {
        "formatVersion": 1,
        "exportType": "archiveData",
        "exportedAt": dt.datetime.now().isoformat(timespec="seconds"),
        "archive": archive_items,
    }


def phone_export_payload(kind: str, selected_locomotives: list[str] | None = None, date_from: str = "", date_to: str = "") -> dict:
    kind = text(kind).strip().lower()
    if kind == "reference":
        return phone_reference_export_payload()
    if kind == "archive":
        return phone_archive_export_payload(selected_locomotives, date_from, date_to)
    raise ValueError("Неизвестный тип экспорта.")


def parse_phone_json_payload(raw: bytes) -> object:
    data = json.loads(raw.decode("utf-8"))
    if isinstance(data, dict) and "payload" in data:
        payload = data.get("payload")
        if isinstance(payload, str):
            return json.loads(payload)
        return payload
    if isinstance(data, str):
        text_data = data.strip()
        if text_data.startswith("{") or text_data.startswith("["):
            try:
                return json.loads(text_data)
            except Exception:
                pass
        try:
            decoded = base64.b64decode(text_data)
            try:
                return json.loads(zlib.decompress(decoded).decode("utf-8"))
            except Exception:
                return json.loads(decoded.decode("utf-8"))
        except Exception:
            return data
    return data


def import_phone_reference_payload(payload: dict) -> dict:
    reference = payload.get("referenceData") if isinstance(payload.get("referenceData"), dict) else payload
    locomotives = reference.get("locomotives") if isinstance(reference, dict) else None
    if not isinstance(locomotives, list):
        return {"error": "Некорректный справочник локомотивов."}, HTTPStatus.BAD_REQUEST

    grouped: dict[str, dict[str, object]] = {}
    for item in locomotives:
        if not isinstance(item, dict):
            continue
        series = text(item.get("series")).strip().upper()
        number = text(item.get("number")).strip()
        if not number:
            continue
        key = f"{series}|{number}"
        candidate = {
            "series": series,
            "number": number,
            "wheelPairCount": parse_excel_int(item.get("wheelPairCount")) or len(item.get("wheelPairs") or []) or locomotive_axis_count(series, number),
            "sortOrder": parse_excel_int(item.get("sortOrder")) or 0,
            "updatedAt": parse_excel_int(item.get("updatedAt")) or 0,
            "deletedAt": parse_excel_int(item.get("deletedAt")) or 0,
            "eightDigitNumber": text(item.get("eightDigitNumber")).strip(),
            "manufactureYear": text(item.get("manufactureYear") or item.get("manufacture_year")).strip(),
            "serviceLife": text(item.get("serviceLife") or item.get("service_life")).strip(),
            "wheelPairs": item.get("wheelPairs") if isinstance(item.get("wheelPairs"), list) else [],
        }
        current = grouped.get(key)
        if current is None or (
            candidate["updatedAt"],
            candidate["deletedAt"],
            candidate["sortOrder"],
        ) > (
            current["updatedAt"],
            current["deletedAt"],
            current["sortOrder"],
        ):
            grouped[key] = candidate

    imported_locomotives = 0
    imported_wheel_pairs = 0
    with DB_LOCK, connect() as conn:
        cur = conn.cursor()
        cur.execute("BEGIN")
        for item in grouped.values():
            series = text(item["series"]).strip().upper()
            number = text(item["number"]).strip()
            wheel_pair_count = int(item["wheelPairCount"] or 0)
            sort_order = int(item["sortOrder"] or 0)
            updated_at = int(item["updatedAt"] or 0)
            deleted_at = int(item["deletedAt"] or 0)
            eight_digit_number = text(item["eightDigitNumber"]).strip()
            manufacture_year = text(item["manufactureYear"]).strip()
            service_life = text(item["serviceLife"]).strip()
            upsert_inventory_locomotive(
                cur,
                series,
                number,
                "",
                wheel_pair_count=wheel_pair_count,
                eight_digit_number=eight_digit_number,
                manufacture_year=manufacture_year,
                service_life=service_life,
                sort_order=sort_order,
                updated_at=updated_at or None,
                deleted_at=deleted_at,
            )
            if deleted_at <= 0:
                cur.execute("DELETE FROM kp_data WHERE locomotive=?", (number,))
            wheel_pairs = item.get("wheelPairs")
            if not isinstance(wheel_pairs, list) or not wheel_pairs:
                wheel_pairs = [{"number": index + 1, "axisNumber": index + 1} for index in range(max(1, wheel_pair_count))]
            for pair in wheel_pairs:
                if not isinstance(pair, dict):
                    continue
                pair_number = parse_excel_int(pair.get("number")) or 0
                if pair_number <= 0:
                    continue
                axis_number = parse_excel_int(pair.get("axisNumber")) or pair_number
                cur.execute(
                    "INSERT OR REPLACE INTO kp_data(locomotive, r, c, v) VALUES(?, ?, ?, ?)",
                    (number, pair_number - 1, 1, str(axis_number)),
                )
                diameter_left = parse_float_value(pair.get("diameterLeft"))
                diameter_right = parse_float_value(pair.get("diameterRight"))
                if diameter_left is not None:
                    cur.execute(
                        "INSERT OR REPLACE INTO kp_data(locomotive, r, c, v) VALUES(?, ?, ?, ?)",
                        (number, pair_number - 1, 2, str(diameter_left)),
                    )
                if diameter_right is not None:
                    cur.execute(
                        "INSERT OR REPLACE INTO kp_data(locomotive, r, c, v) VALUES(?, ?, ?, ?)",
                        (number, pair_number - 1, 3, str(diameter_right)),
                    )
            imported_locomotives += 1
            imported_wheel_pairs += len(wheel_pairs)
        conn.commit()

    return {"ok": True, "imported_locomotives": imported_locomotives, "imported_wheel_pairs": imported_wheel_pairs}


def import_phone_measurement_payload(payload: dict) -> dict:
    measurement = payload.get("measurement") if isinstance(payload.get("measurement"), dict) else payload
    if not isinstance(measurement, dict):
        return {"error": "Некорректный замер."}, HTTPStatus.BAD_REQUEST

    locomotive = measurement.get("locomotive")
    if not isinstance(locomotive, dict):
        return {"error": "Не найден локомотив в замере."}, HTTPStatus.BAD_REQUEST

    series = text(locomotive.get("series")).strip()
    number = text(locomotive.get("number")).strip()
    measurement_date = text(measurement.get("measurementDate")).strip() or dt.date.today().isoformat()
    repair_type = normalize_repair_type(measurement.get("repairType"))
    measurement_id = text(measurement.get("measurementId")).strip() or f"{measurement_date}:{number}:{repair_type}"
    wheel_pairs = measurement.get("wheelPairs")
    if not isinstance(wheel_pairs, list) or not wheel_pairs:
        return {"error": "В замере нет колесных пар."}, HTTPStatus.BAD_REQUEST

    try:
        year = int(measurement_date[:4])
    except Exception:
        year = dt.date.today().year

    wheel_pair_count = parse_excel_int(locomotive.get("wheelPairCount")) or len(wheel_pairs) or locomotive_axis_count(series, number)
    section_count = default_section_count(wheel_pair_count)
    section_sizes: list[int] = []
    base = max(1, wheel_pair_count) // max(1, section_count)
    remainder = max(1, wheel_pair_count) % max(1, section_count)
    for index in range(max(1, section_count)):
        section_sizes.append(base + (1 if index < remainder else 0))

    archive_rows: list[tuple[int, str, str, str, int, int, str]] = []
    imported_cells = 0
    with DB_LOCK, connect() as conn:
        cur = conn.cursor()
        cur.execute("BEGIN")
        upsert_inventory_locomotive(cur, series, number)
        cur.execute(
            "DELETE FROM archive_data WHERE y=? AND measurement_date=? AND locomotive=? AND repair_type=?",
            (year, measurement_date, number, repair_type),
        )
        for pair in wheel_pairs:
            if not isinstance(pair, dict):
                continue
            pair_number = parse_excel_int(pair.get("number")) or 0
            if pair_number <= 0:
                continue
            row_index = pair_number - 1
            running = 0
            section_value = "1"
            for section_index, span in enumerate(section_sizes, start=1):
                running += span
                if row_index < running:
                    section_value = str(section_index)
                    break
            left = pair.get("left") if isinstance(pair.get("left"), dict) else {}
            right = pair.get("right") if isinstance(pair.get("right"), dict) else {}
            values = {
                0: section_value,
                1: str(pair_number),
                2: excel_num_text(left.get("flangeWear")),
                3: excel_num_text(right.get("flangeWear")),
                4: excel_num_text(left.get("flangeThickness")),
                5: excel_num_text(right.get("flangeThickness")),
                6: excel_num_text(left.get("flangeSteepness")),
                7: excel_num_text(right.get("flangeSteepness")),
                8: excel_num_text(left.get("bandageThickness")),
                9: excel_num_text(right.get("bandageThickness")),
                10: excel_num_text(left.get("bandageDiameter")),
                11: excel_num_text(right.get("bandageDiameter")),
            }
            for col, value in values.items():
                value = text(value).strip()
                if value:
                    archive_rows.append((year, measurement_date, number, repair_type, row_index + 2, col, value))
                    imported_cells += 1
        cur.executemany(
            "INSERT OR REPLACE INTO archive_data (y, measurement_date, locomotive, repair_type, r, c, v) VALUES (?, ?, ?, ?, ?, ?, ?)",
            archive_rows,
        )
        conn.commit()

    return {
        "ok": True,
        "imported_measurements": 1,
        "imported_cells": imported_cells,
        "measurement_id": measurement_id,
    }


def import_phone_payload(payload: object) -> dict:
    if isinstance(payload, dict):
        export_type = text(payload.get("exportType")).strip()
        if payload.get("archiveData") is not None or export_type == "archiveData":
            archive = payload.get("archiveData") if isinstance(payload.get("archiveData"), dict) else payload
            archive_items = archive.get("archive") if isinstance(archive, dict) else None
            if isinstance(archive_items, list):
                imported_measurements = 0
                imported_cells = 0
                for item in archive_items:
                    result = import_phone_measurement_payload(item)
                    if isinstance(result, tuple):
                        return result[0], result[1]
                    imported_measurements += int(result.get("imported_measurements", 0))
                    imported_cells += int(result.get("imported_cells", 0))
                return {"ok": True, "imported_measurements": imported_measurements, "imported_cells": imported_cells}
            return import_phone_measurement_payload(archive if isinstance(archive, dict) else payload)
        if payload.get("referenceData") is not None or export_type == "referenceData":
            reference = payload.get("referenceData") if isinstance(payload.get("referenceData"), dict) else payload
            return import_phone_reference_payload(reference)
        if payload.get("measurement") is not None:
            measurement = payload.get("measurement")
            if isinstance(measurement, dict):
                return import_phone_measurement_payload(measurement)
        if "locomotives" in payload:
            return import_phone_reference_payload(payload)
        if "wheelPairs" in payload and "measurementDate" in payload:
            return import_phone_measurement_payload(payload)
    return {"error": "Не удалось распознать телефонные данные."}, HTTPStatus.BAD_REQUEST


def import_archive_excel_bytes(data: bytes) -> dict:
    _, load_workbook, *_ = require_openpyxl()
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active

    aliases = {
        "measurement_date": {"датазамера", "датавыполненияобмера", "датаобмера", "measurementdate"},
        "series": {"серия", "series"},
        "locomotive": {"локомотив", "номерлокомотива", "locomotive"},
        "repair_type": {"видремонта", "видрем", "repairtype"},
        "wheel_pair": {"номеркп", "wheelpair", "wheelpairnumber", "ось", "номероси", "kp"},
        "section": {"секция", "вагон", "section"},
        "prokat_left": {"прокатлев", "левпрокат", "prokatleft", "flangewearleft"},
        "prokat_right": {"прокатправ", "правпрокат", "prokatright", "flangewearright"},
        "greben_left": {"толщинагребнялев", "левтолщинагребня", "grebenleft", "flangethicknessleft"},
        "greben_right": {"толщинагребняправ", "правтолщинагребня", "grebenright", "flangethicknessright"},
        "krut_left": {"крутизнагребнялев", "левкрутизнагребня", "krutleft", "flangesteepnessleft"},
        "krut_right": {"крутизнагребняправ", "правкрутизнагребня", "krutright", "flangesteepnessright"},
        "bandage_thickness_left": {"толщинабандажалева", "толщинабандажалев", "leftbandagethickness", "bandagethicknessleft"},
        "bandage_thickness_right": {"толщинабандажаправ", "rightbandagethickness", "bandagethicknessright"},
        "bandage_diameter_left": {"диаметрбандажалева", "диаметрбандажалев", "leftbandagediameter", "bandagediameterleft"},
        "bandage_diameter_right": {"диаметрбандажаправ", "rightbandagediameter", "bandagediameterright"},
    }

    header_row_index = None
    header_columns: dict[str, int] = {}
    for row_index, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
        normalized = [normalize_excel_header(value) for value in row]
        found: dict[str, int] = {}
        for key, alias_set in aliases.items():
            for idx, header in enumerate(normalized):
                if header in alias_set:
                    found[key] = idx
                    break
        if len(found) >= 5:
            header_row_index = row_index
            header_columns = found
            break

    if header_row_index is None:
        return {"error": "Не найдена строка заголовков."}, HTTPStatus.BAD_REQUEST

    missing_headers = [name for name in ("measurement_date", "locomotive", "repair_type", "wheel_pair") if name not in header_columns]
    if missing_headers:
        return {"error": "В Excel не найдены обязательные колонки: " + ", ".join(missing_headers)}, HTTPStatus.BAD_REQUEST

    required_metric_columns = {
        "prokat_left": 2,
        "prokat_right": 3,
        "greben_left": 4,
        "greben_right": 5,
        "krut_left": 6,
        "krut_right": 7,
        "bandage_thickness_left": 8,
        "bandage_thickness_right": 9,
    }
    optional_metric_columns = {"bandage_diameter_left": 10, "bandage_diameter_right": 11}

    measurements: dict[tuple[str, str, str], dict] = {}
    errors: list[str] = []
    current_series = current_date = current_loco = current_repair = current_section = ""

    def row_value(row, column_name: str) -> str:
        idx = header_columns.get(column_name)
        if idx is None or idx >= len(row):
            return ""
        return excel_cell_text(row[idx])

    for row_index, row in enumerate(ws.iter_rows(min_row=header_row_index + 1, values_only=True), start=header_row_index + 1):
        if not row or all(cell is None or text(cell).strip() == "" for cell in row):
            continue
        row_date = parse_excel_date(row_value(row, "measurement_date")) or current_date
        row_loco = row_value(row, "locomotive") or current_loco
        row_repair = row_value(row, "repair_type") or current_repair
        row_series = row_value(row, "series") or current_series
        row_section = row_value(row, "section") or current_section
        wheel_pair_number = parse_excel_int(row_value(row, "wheel_pair"))

        if row_value(row, "measurement_date"):
            current_date = row_date
        if row_value(row, "locomotive"):
            current_loco = row_loco
        if row_value(row, "repair_type"):
            current_repair = row_repair
        if row_value(row, "series"):
            current_series = row_series
        if row_value(row, "section"):
            current_section = row_section

        if not row_date or not row_loco or not row_repair:
            errors.append(f"Строка {row_index}: не заполнены дата, локомотив или вид ремонта.")
            continue
        if wheel_pair_number is None or wheel_pair_number <= 0:
            errors.append(f"Строка {row_index}: не указан корректный номер КП.")
            continue

        values: dict[int, str] = {}
        missing_fields = []
        for column_name, db_col in required_metric_columns.items():
            value = row_value(row, column_name)
            if value == "":
                missing_fields.append(column_name)
            else:
                values[db_col] = value
        for column_name, db_col in optional_metric_columns.items():
            value = row_value(row, column_name)
            if value != "":
                values[db_col] = value
        if missing_fields:
            errors.append(f"Строка {row_index}: не заполнены поля " + ", ".join(missing_fields))
            continue

        group = measurements.setdefault((row_date, row_loco, row_repair), {"series": row_series, "rows": {}, "sections": {}})
        if row_series and not group["series"]:
            group["series"] = row_series
        if row_section:
            group["sections"][wheel_pair_number] = row_section
        group["rows"][wheel_pair_number] = values

    if not measurements:
        return {"error": "Не удалось импортировать ни одного замера.", "errors": errors[:20]}, HTTPStatus.BAD_REQUEST

    imported_measurements = 0
    imported_cells = 0
    with DB_LOCK, connect() as conn:
        cur = conn.cursor()
        kp_cache: dict[str, dict[tuple[int, int], str]] = {}
        cur.execute("BEGIN")
        for (measurement_date, locomotive, repair_type), meta in sorted(measurements.items()):
            wheel_pair_numbers = sorted(meta["rows"].keys())
            if not wheel_pair_numbers:
                continue
            year = int(measurement_date[:4]) if len(measurement_date) >= 4 and measurement_date[:4].isdigit() else dt.date.today().year
            series = meta["series"] or series_for_locomotive(cur, locomotive)
            ensure_import_locomotive(cur, series, locomotive, max(wheel_pair_numbers))
            if locomotive not in kp_cache:
                kp_rows = cur.execute("SELECT r, c, v FROM kp_data WHERE locomotive=? AND c IN (2, 3)", (locomotive,)).fetchall()
                kp_cache[locomotive] = {(int(row["r"]), int(row["c"])): text(row["v"]).strip() for row in kp_rows}

            axis_count = locomotive_axis_count(series_for_locomotive(cur, locomotive), locomotive)
            db_rows: list[tuple[int, str, str, str, int, int, str]] = []
            for wheel_pair_number in wheel_pair_numbers:
                row_values = meta["rows"][wheel_pair_number]
                table_row = wheel_pair_number + 1
                section_value = meta["sections"].get(wheel_pair_number, "1" if axis_count == 6 else str(((wheel_pair_number - 1) // 4) + 1))
                kp_index = wheel_pair_number - 1
                bandage_left = row_values.get(8, "")
                bandage_right = row_values.get(9, "")
                diameter_left = row_values.get(10, "")
                diameter_right = row_values.get(11, "")
                if not diameter_left:
                    kp_left = parse_float_value(kp_cache[locomotive].get((kp_index, 2), ""))
                    bandage_left_value = parse_float_value(bandage_left)
                    if kp_left is not None and bandage_left_value is not None:
                        diameter_left = str(int(round(kp_left + bandage_left_value * 2)))
                if not diameter_right:
                    kp_right = parse_float_value(kp_cache[locomotive].get((kp_index, 3), ""))
                    bandage_right_value = parse_float_value(bandage_right)
                    if kp_right is not None and bandage_right_value is not None:
                        diameter_right = str(int(round(kp_right + bandage_right_value * 2)))

                full_values = {
                    0: section_value,
                    1: str(wheel_pair_number),
                    2: row_values.get(2, ""),
                    3: row_values.get(3, ""),
                    4: row_values.get(4, ""),
                    5: row_values.get(5, ""),
                    6: row_values.get(6, ""),
                    7: row_values.get(7, ""),
                    8: bandage_left,
                    9: bandage_right,
                    10: diameter_left,
                    11: diameter_right,
                }
                for col, value in full_values.items():
                    if value != "":
                        db_rows.append((year, measurement_date, locomotive, repair_type, table_row, col, value))
            if not db_rows:
                continue
            cur.execute(
                "DELETE FROM archive_data WHERE y=? AND measurement_date=? AND locomotive=? AND repair_type=?",
                (year, measurement_date, locomotive, repair_type),
            )
            cur.executemany(
                "INSERT OR REPLACE INTO archive_data (y, measurement_date, locomotive, repair_type, r, c, v) VALUES (?, ?, ?, ?, ?, ?, ?)",
                db_rows,
            )
            imported_measurements += 1
            imported_cells += len(db_rows)
        conn.commit()

    return {"ok": True, "imported_measurements": imported_measurements, "imported_cells": imported_cells, "errors": errors[:10]}



def build_schedule_eml_bytes(filter_choice: str = "all") -> bytes:
    """Генерация файла черновика письма .eml с графиком замеров колесных пар."""
    import datetime
    year = datetime.date.today().year
    archive_rows = load_archive_rows()

    with DB_LOCK, connect() as conn:
        cur = conn.cursor()

        inv_rows = cur.execute(
            "SELECT ser, num FROM inventory WHERE TRIM(COALESCE(num, '')) <> '' AND COALESCE(deleted_at, 0) = 0"
        ).fetchall()
        inventory_map = {}
        for r in inv_rows:
            num = str(r["num"]).strip()
            ser = str(r["ser"]).strip().upper()
            inventory_map[num] = ser

        def normalize_series(s):
            return re.sub(r'[^a-zA-Z0-9а-яА-Я]', '', s).strip().upper()

        def unit_key_from_cells(s, n):
            series = normalize_series(s)
            if not series:
                series = 'ТЭМ2УМ'
            num = str(n or '').strip()
            if not num:
                return ''
            return f"{series} №{num}"

        repairs_data = cur.execute(
            "SELECT m, r, c, v FROM repairs WHERE y=? AND t='plan' ORDER BY m, r, c",
            (year,)
        ).fetchall()

        month_rows = {}
        for row in repairs_data:
            m = str(row["m"])
            r = int(row["r"])
            c = int(row["c"])
            v = str(row["v"])
            if m not in month_rows:
                month_rows[m] = {}
            if r not in month_rows[m]:
                month_rows[m][r] = {"cells": [""] * 35, "excluded": False}
            if c == -1:
                month_rows[m][r]["excluded"] = True
            elif 0 <= c <= 2:
                month_rows[m][r]["cells"][c] = v
            elif 3 <= c <= 33:
                month_rows[m][r]["cells"][c + 1] = v

        MONTHS_RU = ["Январь", "Февраль", "Март", "Апрель", "Май", "Июнь", "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"]
        month_num_map = {m: i + 1 for i, m in enumerate(MONTHS_RU)}

        units = {}
        excluded_units = set()
        for m_name, rows_dict in month_rows.items():
            m_num = month_num_map.get(m_name, 1)
            for r_idx, row_info in rows_dict.items():
                cells = row_info["cells"]
                s_val = cells[1]
                n_val = cells[2]
                unit_key = unit_key_from_cells(s_val, n_val)
                if not unit_key:
                    continue
                if row_info["excluded"]:
                    excluded_units.add(unit_key)
                if unit_key not in units:
                    units[unit_key] = {"series": s_val or "ТЭМ-2УМ", "number": n_val, "repairs": []}
                
                for col in range(4, 35):
                    cell_val = cells[col]
                    if cell_val:
                        repair_type = cell_val.strip().upper()
                        if repair_type in ('ТО2', 'ТО3', 'ТР1', 'ТР', 'ТР2', 'ТР3', 'СР', 'КР', 'TO2', 'TO3'):
                            day = col - 3
                            try:
                                candidate_date = datetime.date(year, m_num, day)
                                units[unit_key]["repairs"].append({
                                    "date": candidate_date,
                                    "type": repair_type
                                })
                            except ValueError:
                                pass

        for ex in excluded_units:
            if ex in units:
                del units[ex]

        # Последние замеры КП
        latest_by_unit = {}
        for row in archive_rows:
            m_date = row.get("measurement_date")
            if not m_date:
                continue
            loco = str(row.get("locomotive") or "").strip()
            if not loco:
                continue
            ser = inventory_map.get(loco, "ТЭМ-2УМ")
            u_key = unit_key_from_cells(ser, loco)
            if not u_key:
                continue
            try:
                t = datetime.date.fromisoformat(m_date)
                if u_key not in latest_by_unit or t > latest_by_unit[u_key]["date"]:
                    latest_by_unit[u_key] = {"date": t, "date_str": m_date}
            except Exception:
                pass

        KP_RECHECK_DAYS = 30
        today = datetime.date.today()
        best_by_unit = {}

        for u_key, u_data in units.items():
            last_meas = latest_by_unit.get(u_key)
            last_date_str = last_meas["date_str"] if last_meas else "Нет данных"
            limit_date = (last_meas["date"] + datetime.timedelta(days=KP_RECHECK_DAYS)) if last_meas else None
            
            best_repair = None
            for r in u_data["repairs"]:
                r_date = r["date"]
                if last_meas:
                    if last_meas["date"] < r_date <= limit_date:
                        if not best_repair or r_date > best_repair["date"]:
                            best_repair = r
                else:
                    if r_date >= today:
                        if not best_repair or r_date < best_repair["date"]:
                            best_repair = r
            
            best_by_unit[u_key] = {
                "series": u_data["series"],
                "number": u_data["number"],
                "last_date_str": last_date_str,
                "limit_date": limit_date,
                "best_repair": best_repair
            }

        choices_order = [str(r["num"]).strip() for r in inv_rows]
        loco_order_map = {num: idx for idx, num in enumerate(choices_order)}

        table_rows = list(best_by_unit.values())

        if filter_choice == "tem":
            table_rows = [r for r in table_rows if "ТЭМ" in r["series"].upper() or "TEM" in r["series"].upper()]
        elif filter_choice == "pe2m":
            table_rows = [r for r in table_rows if "ПЭ" in r["series"].upper() or "PE" in r["series"].upper()]

        table_rows.sort(key=lambda r: loco_order_map.get(r["number"], 9999))

        rows_html = []
        for r in table_rows:
            limit_str = r["limit_date"].strftime("%d.%m.%Y") if r["limit_date"] else "-"
            last_str = r["last_date_str"]
            if last_str != "Нет данных" and "-" in last_str:
                parts = last_str.split("-")
                if len(parts) == 3:
                    last_str = f"{parts[2]}.{parts[1]}.{parts[0]}"
            
            is_overdue = r["limit_date"] < today if r["limit_date"] else False
            
            if r["best_repair"]:
                best_str = f"{r['best_repair']['date'].strftime('%d.%m.%Y')} ({r['best_repair']['type']})"
            else:
                best_str = "<span style='color:red;'>Нет подходящего ремонта</span>"

            tr_style = "background-color:#ffebee;" if is_overdue else ""
            limit_style = "color:red; font-weight:bold;" if is_overdue else ""

            rows_html.append(f"""
            <tr style='{tr_style}'>
                <td style='border:1px solid #ccc; padding:6px 10px; text-align:center;'>{r['series']}</td>
                <td style='border:1px solid #ccc; padding:6px 10px; text-align:center; font-weight:bold;'>{r['number']}</td>
                <td style='border:1px solid #ccc; padding:6px 10px; text-align:center;'>{last_str}</td>
                <td style='border:1px solid #ccc; padding:6px 10px; text-align:center; {limit_style}'>{limit_str}</td>
                <td style='border:1px solid #ccc; padding:6px 10px; text-align:center;'>{best_str}</td>
            </tr>
            """)

    html_body = f"""<!DOCTYPE html>
    <html>
    <head>
    <meta charset="utf-8">
    <style>
        body {{ font-family: Arial, sans-serif; font-size: 14px; color: #333; }}
        table {{ border-collapse: collapse; width: 100%; max-width: 800px; margin-top: 12px; }}
        th {{ background-color: #276ef1; color: white; border: 1px solid #ccc; padding: 8px 10px; text-align: center; }}
        td {{ border: 1px solid #ccc; padding: 6px 10px; }}
    </style>
    </head>
    <body>
        <h2>График проведения замеров колесных пар локомотивов</h2>
        <p>Сформировано автоматически на основании актуального графика ППР и базы замеров КП.</p>
        <table>
            <thead>
                <tr>
                    <th>Серия</th>
                    <th>Номер</th>
                    <th>Последний замер</th>
                    <th>Крайний срок</th>
                    <th>Следующий по плану</th>
                </tr>
            </thead>
            <tbody>
                {"".join(rows_html)}
            </tbody>
        </table>
        <br>
        <p>Убедительная просьба придерживаться данного графика и ставить локомотивы в депо для проведения замеров.</p>
    </body>
    </html>
    """

    encoded_subject = base64.b64encode("График замеров колесных пар локомотивов".encode('utf-8')).decode('utf-8')
    mime_subject = f"=?utf-8?B?{encoded_subject}?="
    
    eml_headers = [
        "MIME-Version: 1.0",
        "To: TerentevPS@kolagmk.ru; TeterinEYu@kolagmk.ru; StankevichMM@kolagmk.ru; GundorovAO@kolagmk.ru",
        f"Subject: {mime_subject}",
        "X-Unsent: 1",
        "Content-Type: text/html; charset=utf-8",
        "",
        html_body,
    ]
    eml_data = "\r\n".join(eml_headers)
    return eml_data.encode('utf-8')

