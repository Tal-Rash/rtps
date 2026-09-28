# -*- coding: utf-8 -*-
"""
Утилита локальной расшифровки скачанных Excel-файлов (.xlsx).
Заменяет анонимизированные метки (Работник №X, СотрудникПолн №X, Должность №X, ID_00X)
на реальные данные сотрудников из локального файла employees_private.csv / employees_private.xlsx.
"""

import sys
import csv
import openpyxl
from pathlib import Path

def load_mapping():
    csv_file = Path("employees_private.csv")
    xlsx_file = Path("employees_private.xlsx")
    
    mapping = {}
    
    if csv_file.exists():
        with open(csv_file, "r", encoding="utf-8-sig") as f:
            reader = csv.reader(f, delimiter=";")
            for row in reader:
                if not row or len(row) < 2 or row[0].startswith("Код"):
                    continue
                anon_id = row[0].strip()
                pos = row[1].strip() if len(row) > 1 else ""
                full_fio = row[2].strip() if len(row) > 2 else ""
                short_fio = row[3].strip() if len(row) > 3 else full_fio
                tab_num = row[4].strip() if len(row) > 4 else anon_id
                
                num_str = anon_id.replace("ID_", "")
                if num_str.isdigit():
                    num = int(num_str)
                    mapping[f"Работник №{num}"] = short_fio or full_fio
                    mapping[f"СотрудникПолн №{num}"] = full_fio
                    mapping[f"Сотрудник №{num}"] = full_fio
                    if pos:
                        mapping[f"Должность №{num}"] = pos
                    mapping[anon_id] = tab_num
                    
    elif xlsx_file.exists():
        wb = openpyxl.load_workbook(xlsx_file)
        ws = wb.active
        for row in ws.iter_rows(values_only=True):
            if not row or not row[0] or str(row[0]).startswith("Код"):
                continue
            anon_id = str(row[0]).strip()
            pos = str(row[1]).strip() if len(row) > 1 and row[1] else ""
            full_fio = str(row[2]).strip() if len(row) > 2 and row[2] else ""
            short_fio = str(row[3]).strip() if len(row) > 3 and row[3] else full_fio
            tab_num = str(row[4]).strip() if len(row) > 4 and row[4] else anon_id
            
            num_str = anon_id.replace("ID_", "")
            if num_str.isdigit():
                num = int(num_str)
                mapping[f"Работник №{num}"] = short_fio or full_fio
                mapping[f"СотрудникПолн №{num}"] = full_fio
                mapping[f"Сотрудник №{num}"] = full_fio
                if pos:
                    mapping[f"Должность №{num}"] = pos
                mapping[anon_id] = tab_num
                
    return mapping

def process_excel_file(file_path: str, mapping: dict):
    path = Path(file_path)
    if not path.exists():
        print(f"Файл {file_path} не найден!")
        return

    print(f"Обработка файла {path.name}...")
    wb = openpyxl.load_workbook(path)
    count = 0

    sorted_pairs = sorted(mapping.items(), key=lambda x: len(x[0]), reverse=True)

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str):
                    val = cell.value
                    new_val = val
                    for k, v in sorted_pairs:
                        if k in new_val:
                            new_val = new_val.replace(k, str(v))
                            count += 1
                    if new_val != val:
                        cell.value = new_val

    out_path = path.parent / f"{path.stem}_реальный{path.suffix}"
    wb.save(out_path)
    print(f"Успешно заменено элементов: {count}. Сохранен файл: {out_path.name}")

if __name__ == "__main__":
    mapping = load_mapping()
    if not mapping:
        print("Ошибка: Локальный словарь employees_private.csv / employees_private.xlsx не найден!")
    else:
        if len(sys.argv) > 1:
            for f in sys.argv[1:]:
                process_excel_file(f, mapping)
        else:
            file_input = input("Введите путь к скачанному Excel файлу (или перетащите файл сюда): ").strip('"')
            if file_input:
                process_excel_file(file_input, mapping)
